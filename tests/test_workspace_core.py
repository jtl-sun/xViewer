from pathlib import Path
import struct
import zipfile
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.drawing.image import Image as XLImage
from PIL import Image

from mdir import __version__
from mdir.workspace_app import (
    EditableWorkbookModel,
    _parse_clipboard_value,
    _pane_hotkey_action,
    _file_name_matches_filter,
    _search_arrow_target_index,
    _left_file_tab_action,
    _selection_workbook_action,
    _type_filter_accepts,
    _merge_recent_directories,
    _send_paths_to_recycle_bin,
    _centered_geometry,
)
from mdir.theme import effective_theme_name, normalize_theme_mode, palette_for
from mdir.excel_viewer import (
    APP_VERSION as EXCEL_VIEWER_VERSION,
    AxisMetrics,
    EmbeddedImage,
    _load_legacy_xls_image_manifest,
    _merge_embedded_images,
    legacy_xls_pdf_preview_path,
    _sheet_content_extent,
)
from mdir.media_viewer import IMAGE_EXTENSIONS, PDF_EXTENSIONS, _load_source_image, _render_pdf_page, fit_zoom_for_size, media_kind_for_path
from mdir.excel_pdf_preview import (
    excel_pdf_cache_path,
    excel_pdf_cache_hit,
    render_excel_pdf_cached,
    ExcelPdfResult,
)
from mdir.workspace_app import APP_VERSION as WORKSPACE_APP_VERSION
from mdir.links import (
    LinkDefinition,
    expand_link_text,
    load_links,
    load_mdir_links,
    parse_links,
    save_links,
)


class WorkspaceCoreTests(unittest.TestCase):
    def test_version_is_consistent_across_entry_points(self):
        self.assertEqual(WORKSPACE_APP_VERSION, __version__)
        self.assertEqual(EXCEL_VIEWER_VERSION, __version__)

    def test_theme_modes_and_palettes(self):
        self.assertEqual(normalize_theme_mode("light"), "Bright")
        self.assertEqual(normalize_theme_mode("DARK"), "Dark")
        self.assertEqual(normalize_theme_mode("anything-else"), "System")
        self.assertEqual(effective_theme_name("Bright"), "Bright")
        self.assertEqual(effective_theme_name("Dark"), "Dark")
        self.assertEqual(palette_for("Bright")["name"], "Bright")
        self.assertEqual(palette_for("Dark")["name"], "Dark")
        self.assertNotEqual(palette_for("Bright")["window"], palette_for("Dark")["window"])

    def test_main_window_geometry_is_centered_on_primary_screen(self):
        self.assertEqual(_centered_geometry(1920, 1080, 1580, 920), "1580x920+170+80")
        self.assertEqual(_centered_geometry(1366, 768, 1580, 920), "1580x920+0+0")

    def test_xviewer_icon_assets_and_installer_use_new_branding(self):
        project = Path(__file__).resolve().parents[1]
        png = project / "mdir" / "assets" / "xviewer-icon.png"
        ico = project / "mdir" / "assets" / "xviewer.ico"
        self.assertTrue(png.is_file())
        self.assertTrue(ico.is_file())
        with Image.open(png) as image:
            self.assertEqual(image.size, (512, 512))
        installer = (project / "install_windows.ps1").read_text(encoding="utf-8")
        self.assertIn(r'mdir\assets\xviewer.ico', installer)
        self.assertIn('("xviewer-" + $Version + ".ico")', installer)
        self.assertIn('SHChangeNotify', installer)
        self.assertIn('ie4uinit.exe', installer)
        self.assertNotIn(r'mdir\assets\xexcel.ico', installer)

    def test_windows_app_identity_is_xviewer(self):
        project = Path(__file__).resolve().parents[1]
        workspace = (project / "mdir" / "workspace_app.py").read_text(encoding="utf-8")
        excel_viewer = (project / "mdir" / "excel_viewer.py").read_text(encoding="utf-8")
        self.assertIn('SetCurrentProcessExplicitAppUserModelID("jtl-sun.xViewer")', workspace)
        self.assertIn('SetCurrentProcessExplicitAppUserModelID("jtl-sun.xViewer")', excel_viewer)

    def test_recent_directory_history_is_mru_and_deduplicated(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            first = root / "first"
            second = root / "second"
            third = root / "third"
            for path in (first, second, third):
                path.mkdir()
            values = [first, second, first, third]
            recent = _merge_recent_directories(second, values, limit=3)
            self.assertEqual(recent, [second.resolve(), first.resolve(), third.resolve()])

    def test_filename_search_filter_is_case_insensitive_and_supports_multiple_terms(self):
        self.assertTrue(_file_name_matches_filter("010719-9 TJ.xlsx", "010719"))
        self.assertTrue(_file_name_matches_filter("010719-9 TJ.xlsx", "010719 tj"))
        self.assertTrue(_file_name_matches_filter("010719-9 TJ.xlsx", "TJ 9"))
        self.assertTrue(_file_name_matches_filter("ABC.XLSX", "abc"))
        self.assertTrue(_file_name_matches_filter("ABC.XLSX", ""))
        self.assertFalse(_file_name_matches_filter("010719-9 TJ.xlsx", "010620"))
        self.assertFalse(_file_name_matches_filter("010719-9 TJ.xlsx", "010719 win"))

    def test_left_file_list_tab_is_blocked(self):
        self.assertEqual(_left_file_tab_action(), "break")

    def test_search_arrow_leaves_search_and_targets_filtered_rows(self):
        self.assertEqual(_search_arrow_target_index(5, None, 1), 0)
        self.assertEqual(_search_arrow_target_index(5, None, -1), 4)
        self.assertEqual(_search_arrow_target_index(5, 1, 1), 2)
        self.assertEqual(_search_arrow_target_index(5, 3, -1), 2)
        self.assertEqual(_search_arrow_target_index(5, 0, -1), 0)
        self.assertEqual(_search_arrow_target_index(5, 4, 1), 4)
        self.assertIsNone(_search_arrow_target_index(0, None, 1))

    def test_single_left_selection_decides_right_workbook_action(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            folder = root / "folder"
            folder.mkdir()
            excel = root / "book.xlsx"
            excel.write_bytes(b"placeholder")
            pdf = root / "spec.pdf"
            pdf.write_bytes(b"placeholder")
            image = root / "photo.png"
            image.write_bytes(b"placeholder")
            text = root / "notes.txt"
            text.write_text("hello", encoding="utf-8")

            self.assertEqual(_selection_workbook_action([excel]), "load")
            self.assertEqual(_selection_workbook_action([pdf]), "load_pdf")
            self.assertEqual(_selection_workbook_action([image]), "load_image")
            self.assertEqual(_selection_workbook_action([folder]), "clear")
            self.assertEqual(_selection_workbook_action([text]), "load_text")
            self.assertEqual(_selection_workbook_action([]), "keep")
            self.assertEqual(_selection_workbook_action([excel, text]), "keep")

    def test_excel_pdf_image_type_filters(self):
        folder = Path("folder")
        self.assertTrue(_type_filter_accepts(folder, True, excel_only=True, pdf_only=False, images_only=False))
        self.assertTrue(_type_filter_accepts("book.xlsx", False, excel_only=True, pdf_only=False, images_only=False))
        self.assertFalse(_type_filter_accepts("spec.pdf", False, excel_only=True, pdf_only=False, images_only=False))
        self.assertTrue(_type_filter_accepts("spec.pdf", False, excel_only=False, pdf_only=True, images_only=False))
        self.assertTrue(_type_filter_accepts("photo.jpg", False, excel_only=False, pdf_only=False, images_only=True))
        self.assertFalse(_type_filter_accepts("notes.txt", False, excel_only=False, pdf_only=False, images_only=True))
        self.assertTrue(_type_filter_accepts("notes.txt", False, excel_only=False, pdf_only=False, images_only=False))

    def test_media_extensions_are_recognized(self):
        self.assertEqual(media_kind_for_path("drawing.PDF"), "pdf")
        self.assertEqual(media_kind_for_path("photo.JPEG"), "image")
        self.assertEqual(media_kind_for_path("photo.webp"), "image")
        self.assertIsNone(media_kind_for_path("notes.txt"))
        self.assertIn(".pdf", PDF_EXTENSIONS)
        self.assertIn(".png", IMAGE_EXTENSIONS)

    def test_image_fit_zoom_keeps_small_images_1_to_1_and_shrinks_large_images(self):
        self.assertEqual(fit_zoom_for_size(400, 300, 1000, 800), 1.0)
        zoom = fit_zoom_for_size(2000, 1000, 1000, 800)
        self.assertGreater(zoom, 0.45)
        self.assertLess(zoom, 0.50)
        tall_zoom = fit_zoom_for_size(500, 2000, 1000, 800)
        self.assertGreater(tall_zoom, 0.35)
        self.assertLess(tall_zoom, 0.40)

    def test_image_and_pdf_render_helpers(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            image_path = root / "photo.png"
            Image.new("RGB", (120, 80), "white").save(image_path)
            loaded = _load_source_image(image_path)
            self.assertEqual(loaded.size, (120, 80))

            pdf_path = root / "sample.pdf"
            Image.new("RGB", (240, 160), "white").save(pdf_path, "PDF")
            rendered, count = _render_pdf_page(pdf_path, 0, 1.0)
            self.assertEqual(count, 1)
            self.assertGreater(rendered.width, 200)
            self.assertGreater(rendered.height, 140)

    def test_excel_pdf_cache_uses_source_metadata_and_reuses_existing_preview(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            source = root / "sample.xlsx"
            source.write_bytes(b"first")
            with patch("mdir.excel_pdf_preview._cache_dir", return_value=root / "cache"):
                first = excel_pdf_cache_path(source)
                first.parent.mkdir(parents=True, exist_ok=True)
                Image.new('RGB', (100, 80), 'white').save(first, 'PDF')
                self.assertEqual(excel_pdf_cache_hit(source), first)
                result = render_excel_pdf_cached(source)
                self.assertTrue(result.ok)
                self.assertTrue(result.from_cache)
                self.assertEqual(result.backend, "cache")

                # Changing source size/mtime produces a new cache identity.
                source.write_bytes(b"second-version-with-different-size")
                second = excel_pdf_cache_path(source)
                self.assertNotEqual(first, second)
                self.assertIsNone(excel_pdf_cache_hit(source))

    def test_excel_open_path_prefers_cached_pdf_engine_not_native_grid(self):
        import inspect
        from mdir.workspace_app import ExcelWorkspaceApp
        source = inspect.getsource(ExcelWorkspaceApp.open_workbook)
        self.assertIn("excel_pdf_cache_hit", source)
        self.assertIn("self._excel_pdf_engine.request", source)
        self.assertIn('self.viewer_kind = "excel_pdf"', source)
        self.assertNotIn("EditableWorkbookModel(path)", source)

    def test_excel_pdf_preview_does_not_prefetch_adjacent_files(self):
        from mdir.excel_pdf_preview import ExcelPdfPreviewEngine
        engine = ExcelPdfPreviewEngine()
        try:
            with patch.object(engine, 'request') as request:
                engine.prefetch([Path('neighbor.xlsx')])
                request.assert_not_called()
        finally:
            engine.close()
            engine._thread.join(2)

    def test_windows_pywin32_dependency_is_declared_for_persistent_excel_engine(self):
        project = Path(__file__).resolve().parents[1]
        pyproject = (project / "pyproject.toml").read_text(encoding="utf-8")
        self.assertIn("pywin32>=306; platform_system == 'Windows'", pyproject)

    def test_parse_clipboard_value(self):
        self.assertEqual(_parse_clipboard_value("12"), 12)
        self.assertEqual(_parse_clipboard_value("12.5"), 12.5)
        self.assertEqual(_parse_clipboard_value("=A1+B1"), "=A1+B1")
        self.assertEqual(_parse_clipboard_value("00123"), "00123")



    def test_link_parser_keeps_mdir_types_and_pane(self):
        values = [
            {"label": "CHOIS", "type": "folder", "target": r"D:\\sys_back\\chois", "pane": "left"},
            {"label": "GitHub", "type": "web", "target": "https://example.com"},
            {"label": "CMD", "type": "action", "target": "powershell_here", "pane": "active"},
            {"label": "Tool", "type": "program", "target": "tool.exe", "args": ["-x"]},
        ]
        links = parse_links(values)
        self.assertEqual([item.kind for item in links], ["folder", "web", "action", "program"])
        self.assertEqual(links[0].pane, "left")
        self.assertEqual(links[3].args, ("-x",))

    def test_links_roundtrip_mdir_compatible_schema(self):
        with TemporaryDirectory() as td:
            path = Path(td) / "links.json"
            links = [
                LinkDefinition("Home", "folder", "{home}"),
                LinkDefinition("GitHub", "web", "https://example.com"),
                LinkDefinition("CMD", "action", "powershell_here"),
                LinkDefinition("Tool", "program", "tool.exe", ("-x",), "right"),
            ]
            save_links(links, path)
            self.assertEqual(load_links(path, Path(td) / "missing-mdir.json"), links)

    def test_load_mdir_keeps_non_folder_items(self):
        with TemporaryDirectory() as td:
            path = Path(td) / "mdir-links.json"
            path.write_text(
                '[{"label":"Work","type":"folder","target":"C:/Work"},'
                '{"label":"GitHub","type":"web","target":"https://example.com"},'
                '{"label":"CMD","type":"action","target":"powershell_here"}]',
                encoding="utf-8",
            )
            links = load_mdir_links(path)
            self.assertEqual([item.kind for item in links], ["folder", "web", "action"])

    def test_legacy_folder_config_migrates_to_full_mdir_list(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            xexcel = root / "xexcel.json"
            mdir = root / "mdir.json"
            xexcel.write_text(
                '[{"label":"Home","target":"{home}"},'
                '{"label":"SPEC","target":"C:/SPEC"}]',
                encoding="utf-8",
            )
            mdir.write_text(
                '[{"label":"Home","type":"folder","target":"{home}"},'
                '{"label":"SPEC","type":"folder","target":"C:/SPEC"},'
                '{"label":"GitHub","type":"web","target":"https://example.com"},'
                '{"label":"CMD","type":"action","target":"powershell_here"}]',
                encoding="utf-8",
            )
            links = load_links(xexcel, mdir)
            self.assertEqual([x.label for x in links], ["Home", "SPEC", "GitHub", "CMD"])

    def test_expand_link_text_supports_mdir_tokens(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            current = root / "current"
            selected = current / "book.xlsx"
            workbook = root / "wb" / "open.xlsx"
            expanded = expand_link_text(
                "{current}|{selected}|{right_selected}|{home}|{project}",
                current=current, selected=selected, workbook=workbook, project=root / "project",
            )
            self.assertIn(str(current), expanded)
            self.assertIn(str(selected), expanded)
            self.assertIn(str(workbook), expanded)
            self.assertIn(str(root / "project"), expanded)

    def test_pane_hotkey_action_windows_and_tk_masks(self):
        self.assertEqual(_pane_hotkey_action("Left", 37, 0x0004), "left")
        self.assertEqual(_pane_hotkey_action("Right", 39, 0x0004), "right")
        self.assertEqual(_pane_hotkey_action("1", 49, 0x20000), "left")
        self.assertEqual(_pane_hotkey_action("2", 50, 0x20000), "right")
        self.assertEqual(_pane_hotkey_action("Left", 37, 0, ctrl_down=True), "left")
        self.assertEqual(_pane_hotkey_action("2", 50, 0, alt_down=True), "right")
        self.assertIsNone(_pane_hotkey_action("Left", 37, 0))
        # On Windows a stale Tk Alt-like state bit must not turn ordinary
        # number input into Alt+1 / Alt+2 when the physical Alt key is up.
        self.assertIsNone(_pane_hotkey_action("1", 49, 0x20000, alt_down=False))
        self.assertIsNone(_pane_hotkey_action("2", 50, 0x20000, alt_down=False))
        self.assertIsNone(_pane_hotkey_action("3", 51, 0x20000, alt_down=False))
        # Without an explicit physical-state override (e.g. X11/Linux), the
        # normal Tk Alt mask remains supported.
        self.assertEqual(_pane_hotkey_action("1", 49, 0x20000), "left")

    def test_edit_and_save_copy(self):
        with TemporaryDirectory() as td:
            source = Path(td) / "source.xlsx"
            target = Path(td) / "target.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "Original"
            wb.save(source)

            model = EditableWorkbookModel(source)
            self.assertTrue(model.editable)
            sheet = model.sheets[0]
            sheet.set_value(1, 1, "Changed")
            sheet.set_value(2, 2, 123)
            self.assertTrue(model.dirty)
            model.save(target, create_backup=False)
            model.close()

            check = load_workbook(target, data_only=False)
            self.assertEqual(check.active["A1"].value, "Changed")
            self.assertEqual(check.active["B2"].value, 123)
            check.close()

    def test_malformed_emf_image_does_not_abort_modern_workbook_load(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            good1 = root / "good1.png"
            good2 = root / "good2.png"
            Image.new("RGB", (32, 24), (20, 100, 180)).save(good1)
            Image.new("RGB", (28, 20), (180, 100, 20)).save(good2)

            source = root / "source.xlsx"
            broken = root / "broken-image.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "still opens"
            ws.add_image(XLImage(str(good1)), "B2")
            ws.add_image(XLImage(str(good2)), "F2")
            wb.save(source)

            # Build a minimal EMF header with a valid bbox but zero frame
            # dimensions. Pillow's WMF/EMF parser divides by the frame width
            # and used to raise ZeroDivisionError while openpyxl loaded the
            # workbook drawing layer.
            malformed = bytearray(44)
            malformed[0:4] = b"\x01\x00\x00\x00"
            struct.pack_into("<iiii", malformed, 8, 0, 0, 100, 100)
            struct.pack_into("<iiii", malformed, 24, 0, 0, 0, 0)
            malformed[40:44] = b" EMF"

            with zipfile.ZipFile(source, "r") as zin:
                media = sorted(name for name in zin.namelist() if name.startswith("xl/media/"))
                self.assertEqual(len(media), 2)
                bad_name = media[-1]
                with zipfile.ZipFile(broken, "w", zipfile.ZIP_DEFLATED) as zout:
                    for info in zin.infolist():
                        payload = bytes(malformed) if info.filename == bad_name else zin.read(info.filename)
                        zout.writestr(info, payload)

            model = EditableWorkbookModel(broken)
            try:
                self.assertEqual(model.sheets[0].value(1, 1), "still opens")
                self.assertFalse(model.editable)
                self.assertGreaterEqual(len(model.skipped_image_errors), 1)
                self.assertEqual(len(model.sheets[0].images), 1)
                self.assertIn("Safe View", model.kind)
            finally:
                model.close()

    def test_missing_null_drawing_part_does_not_abort_workbook_load(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            image1 = root / "sheet1.png"
            image2 = root / "sheet2.png"
            Image.new("RGB", (24, 18), (20, 120, 60)).save(image1)
            Image.new("RGB", (26, 20), (160, 80, 30)).save(image2)

            source = root / "source.xlsx"
            broken = root / "missing-drawing-part.xlsx"
            wb = Workbook()
            ws1 = wb.active
            ws1.title = "BrokenDrawing"
            ws1["A1"] = "sheet one still opens"
            ws1.add_image(XLImage(str(image1)), "B2")
            ws2 = wb.create_sheet("GoodDrawing")
            ws2["A1"] = "sheet two still opens"
            ws2.add_image(XLImage(str(image2)), "C3")
            wb.save(source)

            # Reproduce workbooks that contain a stale drawing relationship
            # whose target is literally /xl/drawings/NULL. openpyxl normally
            # raises KeyError: There is no item named 'xl/drawings/NULL'.
            with zipfile.ZipFile(source, "r") as zin:
                with zipfile.ZipFile(broken, "w", zipfile.ZIP_DEFLATED) as zout:
                    for info in zin.infolist():
                        payload = zin.read(info.filename)
                        if info.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                            payload = payload.replace(
                                b"/xl/drawings/drawing1.xml",
                                b"/xl/drawings/NULL",
                            )
                        zout.writestr(info, payload)

            model = EditableWorkbookModel(broken)
            try:
                self.assertEqual(model.sheets[0].value(1, 1), "sheet one still opens")
                self.assertEqual(model.sheets[1].value(1, 1), "sheet two still opens")
                self.assertFalse(model.editable)
                self.assertGreaterEqual(len(model.skipped_image_errors), 1)
                self.assertTrue(any("NULL" in item for item in model.skipped_image_errors))
                self.assertEqual(len(model.sheets[0].images), 0)
                self.assertEqual(len(model.sheets[1].images), 1)
                self.assertIn("Safe View", model.kind)
            finally:
                model.close()

    def test_embedded_image_survives_edit_and_save(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            image_path = root / "product.png"
            Image.new("RGB", (40, 30), (120, 80, 40)).save(image_path)
            source = root / "image-source.xlsx"
            target = root / "image-target.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "Product"
            ws.add_image(XLImage(str(image_path)), "B2")
            wb.save(source)

            model = EditableWorkbookModel(source)
            self.assertEqual(len(model.sheets[0].images), 1)
            model.sheets[0].set_value(1, 1, "Changed Product")
            model.save(target, create_backup=False)
            model.close()

            check = load_workbook(target)
            self.assertEqual(check.active["A1"].value, "Changed Product")
            self.assertEqual(len(check.active._images), 1)
            check.close()


    def test_legacy_xls_converted_preview_keeps_embedded_images_read_only(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            legacy = root / "legacy.xls"
            legacy.write_bytes(b"legacy-placeholder")
            image_path = root / "legacy-product.png"
            Image.new("RGB", (80, 60), (140, 90, 50)).save(image_path)
            converted = root / "legacy-preview.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "Legacy product"
            ws.add_image(XLImage(str(image_path)), "B2")
            wb.save(converted)

            with patch("mdir.workspace_app.legacy_xls_preview_path", return_value=converted), patch(
                "mdir.workspace_app.legacy_xls_images", return_value={}
            ):
                model = EditableWorkbookModel(legacy)
            try:
                self.assertFalse(model.editable)
                self.assertIn("converted preview", model.kind)
                self.assertEqual(model.sheets[0].text(1, 1), "Legacy product")
                self.assertEqual(len(model.sheets[0].images), 1)
                self.assertGreater(model.sheets[0].images[0].width, 0)
                self.assertGreater(model.sheets[0].images[0].height, 0)
            finally:
                model.close()

    def test_legacy_xls_direct_image_manifest_loads_png_and_position(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            legacy = root / "legacy.xls"
            legacy.write_bytes(b"legacy-placeholder")
            image_dir = root / "legacy-images"
            image_dir.mkdir()
            picture = image_dir / "s001_i0001.png"
            Image.new("RGB", (120, 80), (20, 80, 140)).save(picture)
            manifest = root / "legacy-images.json"
            manifest.write_text(
                '{"version":1,"records":[{"sheet":"Sheet1","row":4,"column":9,'
                '"x_offset":3,"y_offset":5,"width":120,"height":80,'
                '"file":"s001_i0001.png","name":"Picture 1"}]}',
                encoding="utf-8",
            )

            with patch("mdir.excel_viewer._legacy_xls_image_manifest_path", return_value=manifest), patch(
                "mdir.excel_viewer._legacy_xls_image_dir", return_value=image_dir
            ):
                loaded = _load_legacy_xls_image_manifest(legacy)

            self.assertEqual(list(loaded), ["Sheet1"])
            item = loaded["Sheet1"][0]
            self.assertEqual((item.row, item.column), (4, 9))
            self.assertEqual((item.x_offset, item.y_offset), (3, 5))
            self.assertEqual((item.width, item.height), (120, 80))
            self.assertTrue(item.raw.startswith(b"\x89PNG"))

    def test_legacy_xls_direct_images_merge_without_duplicate(self):
        original = EmbeddedImage(row=5, column=10, width=300, height=200, raw=b"one")
        duplicate = EmbeddedImage(row=5, column=10, width=302, height=198, raw=b"two")
        missing = EmbeddedImage(row=20, column=3, width=140, height=90, raw=b"three")
        merged = _merge_embedded_images([original], [duplicate, missing])
        self.assertEqual(len(merged), 2)
        self.assertEqual(merged[0].raw, b"one")
        self.assertEqual(merged[1].raw, b"three")

    def test_legacy_xls_exact_pdf_preview_is_cached(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            legacy = root / "legacy.xls"
            legacy.write_bytes(b"legacy-placeholder")
            rendered = root / "legacy-exact.pdf"

            def fake_render(_source, target, timeout=120):
                Path(target).write_bytes(b"%PDF-1.4\n" + b"x" * 256)
                return True

            with patch("mdir.excel_viewer._legacy_xls_pdf_path", return_value=rendered), patch(
                "mdir.excel_viewer._render_xls_with_excel_pdf", side_effect=fake_render
            ) as renderer:
                first = legacy_xls_pdf_preview_path(legacy)
                second = legacy_xls_pdf_preview_path(legacy)

            self.assertEqual(first, rendered)
            self.assertEqual(second, rendered)
            self.assertEqual(renderer.call_count, 1)
            self.assertGreater(rendered.stat().st_size, 100)

    def test_backup_on_original_overwrite(self):
        with TemporaryDirectory() as td:
            source = Path(td) / "source.xlsx"
            wb = Workbook()
            wb.active["A1"] = "Before"
            wb.save(source)
            model = EditableWorkbookModel(source)
            model.sheets[0].set_value(1, 1, "After")
            model.save()
            backups = list((Path(td) / "xExcel_Backup").glob("source-*.xlsx"))
            self.assertEqual(len(backups), 1)
            model.close()


    def test_excel_layout_fidelity_metadata(self):
        with TemporaryDirectory() as td:
            source = Path(td) / "layout.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.sheet_view.showGridLines = False
            ws.sheet_format.defaultRowHeight = 15
            ws.column_dimensions["B"].hidden = True
            ws.column_dimensions["C"].hidden = True
            ws.column_dimensions["D"].width = 20
            ws.row_dimensions[2].hidden = True
            ws.row_dimensions[4].height = 30
            ws.merge_cells("A1:D1")
            ws["A1"] = "ITEM#"
            ws["A1"].font = Font(name="Calibri", size=16, bold=True)
            ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
            ws["A1"].fill = PatternFill("solid", fgColor="FFF2CC")
            thin = Side(style="thin", color="000000")
            ws["A5"] = "Bordered"
            ws["A5"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
            wb.save(source)

            model = EditableWorkbookModel(source)
            sheet = model.sheets[0]
            cols = sheet.column_overrides(1.0)
            rows = sheet.row_overrides(1.0)
            self.assertFalse(sheet.show_gridlines)
            self.assertEqual(cols[2], 0)
            self.assertEqual(cols[3], 0)
            self.assertGreater(cols[4], 100)
            self.assertEqual(rows[2], 0)
            self.assertGreater(rows[4], 30)
            style = sheet.style(1, 1)
            self.assertTrue(style.bold)
            self.assertEqual(style.font_name, "Calibri")
            self.assertEqual(style.vertical, "center")
            self.assertEqual(style.fill.upper(), "#FFF2CC")
            bordered = sheet.style(5, 1)
            self.assertEqual(bordered.left.style, "thin")
            self.assertEqual(sheet.merged_ranges[0], (1, 1, 4, 1))
            model.close()

    def test_scroll_extent_includes_image_below_used_cells(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            image_path = root / "tall.png"
            Image.new("RGB", (120, 700), (180, 180, 180)).save(image_path)
            source = root / "tall-image.xlsx"
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "Only used cell"
            picture = XLImage(str(image_path))
            picture.width = 120
            picture.height = 700
            ws.add_image(picture, "B2")
            wb.save(source)

            model = EditableWorkbookModel(source)
            sheet = model.sheets[0]
            rows = AxisMetrics(sheet.max_row, sheet.default_row_size(1.0), sheet.row_overrides(1.0))
            cols = AxisMetrics(sheet.max_column, sheet.default_column_size(1.0), sheet.column_overrides(1.0))
            width, height = _sheet_content_extent(sheet, rows, cols, 1.0)
            self.assertGreater(height, rows.total + 500)
            self.assertGreaterEqual(width, cols.total)
            model.close()

    def test_all_worksheet_names_are_loaded(self):
        with TemporaryDirectory() as td:
            source = Path(td) / "many-sheets.xlsx"
            wb = Workbook()
            wb.active.title = "Overview"
            for name in ("Costing", "Images", "Buyer Notes", "Order", "Archive"):
                wb.create_sheet(name)
            wb.save(source)
            model = EditableWorkbookModel(source)
            self.assertEqual(
                [sheet.title for sheet in model.sheets],
                ["Overview", "Costing", "Images", "Buyer Notes", "Order", "Archive"],
            )
            model.close()

    def test_enter_binding_opens_default_application(self):
        import inspect
        from mdir.workspace_app import ExcelWorkspaceApp
        source = inspect.getsource(ExcelWorkspaceApp._file_enter_open_default)
        self.assertIn("_open_path_with_default_application", source)
        self.assertIn("path.is_dir()", source)

    def test_left_double_click_opens_default_application_and_folders_navigate(self):
        import inspect
        from mdir.workspace_app import ExcelWorkspaceApp
        build = inspect.getsource(ExcelWorkspaceApp._build_ui)
        activated = inspect.getsource(ExcelWorkspaceApp._file_activated)
        self.assertIn('bind("<Double-Button-1>", self._file_activated)', build)
        self.assertIn("_open_path_with_default_application", activated)
        self.assertIn("path.is_dir()", activated)
        self.assertNotIn("self.open_workbook(path)", activated)
        self.assertIn("identify_row", activated)

    def test_multiselect_and_delete_binding_are_enabled(self):
        import inspect
        from mdir.workspace_app import ExcelWorkspaceApp
        source = inspect.getsource(ExcelWorkspaceApp._build_ui)
        self.assertIn('selectmode="extended"', source)
        self.assertIn('bind("<Delete>", self._delete_selected_items)', source)

    def test_recycle_helper_never_permanently_deletes(self):
        with TemporaryDirectory() as td:
            root = Path(td)
            one = root / "one.xlsx"
            two = root / "two.xlsx"
            one.write_text("1", encoding="utf-8")
            two.write_text("2", encoding="utf-8")
            captured = []
            with patch("mdir.workspace_app.send2trash", side_effect=lambda p: captured.append(Path(p))):
                moved, failed = _send_paths_to_recycle_bin([one, two])
            self.assertEqual(moved, [one, two])
            self.assertEqual(failed, [])
            self.assertEqual(captured, [one, two])
            # The patched recycle function did not remove the files; this also
            # proves our helper does not fall back to unlink/rmtree itself.
            self.assertTrue(one.exists())
            self.assertTrue(two.exists())


    def test_right_mouse_drag_selection_bindings_are_enabled(self):
        import inspect
        from mdir.workspace_app import ExcelWorkspaceApp
        build = inspect.getsource(ExcelWorkspaceApp._build_ui)
        self.assertIn('bind("<ButtonPress-3>", self._file_right_drag_start)', build)
        self.assertIn('bind("<B3-Motion>", self._file_right_drag_motion)', build)
        self.assertIn('bind("<ButtonRelease-3>", self._file_right_drag_end)', build)
        process = inspect.getsource(ExcelWorkspaceApp._file_right_drag_process_y)
        self.assertIn('_right_drag_seen', inspect.getsource(ExcelWorkspaceApp._file_right_drag_toggle_iid))
        self.assertIn('range(previous + step, index + step, step)', process)


    def test_right_drag_keeps_starting_selected_row_and_adds_crossed_rows(self):
        from mdir.workspace_app import ExcelWorkspaceApp

        class FakeTree:
            def __init__(self):
                self.selected = {"row2"}
                self.focused = None
            def selection(self):
                return tuple(self.selected)
            def selection_add(self, iid):
                self.selected.add(iid)
            def selection_remove(self, iid):
                self.selected.discard(iid)
            def focus(self, iid):
                self.focused = iid
            def see(self, iid):
                pass

        app = ExcelWorkspaceApp.__new__(ExcelWorkspaceApp)
        app.file_tree = FakeTree()
        app._right_drag_seen = set()

        # The bug was that row2 was already selected and got toggled OFF when
        # the right-drag started on it. It must stay selected now.
        app._file_right_drag_toggle_iid("row2")
        app._file_right_drag_toggle_iid("row3")
        app._file_right_drag_toggle_iid("row4")
        self.assertEqual(app.file_tree.selected, {"row2", "row3", "row4"})
        self.assertEqual(app.file_tree.focused, "row4")


if __name__ == "__main__":
    unittest.main()
