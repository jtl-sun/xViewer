from __future__ import annotations

import bisect
import hashlib
import io
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
import threading
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter import font as tkfont

from PIL import Image, ImageTk

from .theme import LIGHT_PALETTE

EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"}
MODERN_EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
APP_TITLE = "xViewer"
from . import __version__ as APP_VERSION

DEFAULT_ROW_HEIGHT = 20
DEFAULT_COLUMN_WIDTH = 64
MIN_ROW_HEIGHT = 14
MAX_ROW_HEIGHT = 240
MIN_COLUMN_WIDTH = 42
MAX_COLUMN_WIDTH = 600
GRID_COLOR = "#d9d9d9"
HEADER_BACKGROUND = "#f1f1f1"
HEADER_FOREGROUND = "#222222"
SELECTION_COLOR = "#217346"
ACTIVE_CELL_FILL = "#e8f3ec"

_LEGACY_XLS_CONVERSION_LOCK = threading.RLock()

_OPENPYXL_SAFE_LOAD_LOCK = threading.RLock()


def _safe_openpyxl_load_workbook(path: Path, **kwargs):
    """Load a modern Excel workbook while isolating malformed embedded images.

    Some legacy-created .xlsx files contain EMF/WMF records whose frame metadata
    is invalid (for example a zero frame width/height). Pillow may then raise
    ``ZeroDivisionError`` while openpyxl is merely scanning drawing images.
    openpyxl normally catches ``OSError`` only, so one malformed picture can make
    the entire workbook fail to open.

    xViewer temporarily wraps openpyxl's per-image constructor so any image decode
    failure is converted to ``OSError``. openpyxl will then skip only that picture
    and continue loading the workbook. If a drawing-level failure still escapes,
    xViewer retries once with drawing images/charts suppressed. The caller receives
    a list of skipped-image diagnostics and should treat the workbook as read-only
    to avoid saving a file after openpyxl has omitted unsupported drawing content.
    """
    from openpyxl import load_workbook
    import openpyxl.reader.drawings as drawings_reader
    import openpyxl.reader.excel as excel_reader

    source = Path(path)
    if source.suffix.lower() not in MODERN_EXCEL_EXTENSIONS:
        # Some suppliers name an OOXML workbook .xls. A binary stream avoids
        # openpyxl's extension check without renaming/modifying the source.
        source = io.BytesIO(source.read_bytes())
    skipped: list[str] = []

    with _OPENPYXL_SAFE_LOAD_LOCK:
        original_image = drawings_reader.Image
        original_find_images = excel_reader.find_images

        def guarded_image(image_source):
            try:
                return original_image(image_source)
            except Exception as exc:
                skipped.append(f"{type(exc).__name__}: {exc}")
                # openpyxl.find_images already knows how to skip OSError.
                raise OSError(f"xViewer skipped malformed embedded image: {exc}") from exc

        def guarded_find_images(archive, drawing_path):
            try:
                return original_find_images(archive, drawing_path)
            except (KeyError, OSError, ZeroDivisionError, ValueError, SyntaxError) as exc:
                # Some producer applications leave a worksheet drawing
                # relationship behind but point it at a non-existent part such
                # as ``xl/drawings/NULL``.  openpyxl raises KeyError before it
                # ever reaches the per-image constructor.  Skip only that
                # broken drawing relationship so other sheets/drawings remain
                # available.
                skipped.append(
                    f"Drawing skipped ({drawing_path!s}): {type(exc).__name__}: {exc}"
                )
                return [], []

        drawings_reader.Image = guarded_image
        excel_reader.find_images = guarded_find_images
        try:
            try:
                workbook = load_workbook(source, **kwargs)
            except (KeyError, OSError, ZeroDivisionError, ValueError, SyntaxError) as exc:
                # Last-resort safe-view fallback: malformed DrawingML, missing
                # drawing package parts, or image metadata can fail before the
                # per-drawing/per-image guards are reached. Load cells/sheets
                # without drawings so the workbook remains viewable.
                skipped.append(f"Drawing fallback: {type(exc).__name__}: {exc}")

                def no_drawings(_archive, _path):
                    return [], []

                excel_reader.find_images = no_drawings
                workbook = load_workbook(source, **kwargs)
        finally:
            drawings_reader.Image = original_image
            excel_reader.find_images = original_find_images

    return workbook, skipped
_LEGACY_XLS_CACHE_VERSION = "2.6.5-excel-script-execution"


def _powershell_script_command(powershell: str, script: str) -> list[str]:
    # -Command - parses stdin interactively. Multiline try/finally and function
    # blocks can remain unexecuted at EOF even with a successful exit status.
    # Pass the complete script as one argument; source paths stay in env vars.
    # STA also supplies the apartment required by Excel's clipboard export.
    return [powershell, "-NoProfile", "-NonInteractive", "-STA",
            "-ExecutionPolicy", "Bypass", "-Command", script]


def _legacy_xls_cache_dir() -> Path:
    """Return the persistent cache directory used for read-only .xls previews."""
    base = os.environ.get("LOCALAPPDATA") or tempfile.gettempdir()
    return Path(base) / "xViewer" / "cache" / "legacy-xls"


def _legacy_xls_cache_key(source: Path) -> str:
    """Create a cache key that changes with source contents and cache format."""
    source = Path(source)
    try:
        resolved = source.resolve()
    except Exception:
        resolved = source.absolute()
    try:
        stat = source.stat()
        stamp = f"{stat.st_mtime_ns}|{stat.st_size}"
    except OSError:
        stamp = "missing"
    token = f"{_LEGACY_XLS_CACHE_VERSION}|{os.path.normcase(str(resolved))}|{stamp}"
    return hashlib.sha256(token.encode("utf-8", "surrogatepass")).hexdigest()[:24]


def _legacy_xls_cache_path(source: Path) -> Path:
    return _legacy_xls_cache_dir() / f"{_legacy_xls_cache_key(source)}.xlsx"


def _legacy_xls_image_dir(source: Path) -> Path:
    return _legacy_xls_cache_dir() / f"{_legacy_xls_cache_key(source)}-images"


def _legacy_xls_image_manifest_path(source: Path) -> Path:
    return _legacy_xls_cache_dir() / f"{_legacy_xls_cache_key(source)}-images.json"


def _legacy_xls_pdf_path(source: Path) -> Path:
    return _legacy_xls_cache_dir() / f"{_legacy_xls_cache_key(source)}-exact.pdf"


def _render_xls_with_excel_pdf(source: Path, target: Path, timeout: int = 120) -> bool:
    """Render a legacy .xls workbook with Microsoft Excel's own engine.

    Very old BIFF drawings can disappear during SaveAs/openpyxl conversion and
    can also fail shape-by-shape clipboard export. Excel's fixed-format renderer
    still draws those objects, so this is the reliable read-only visual fallback.
    """
    if os.name != "nt":
        return False
    powershell = shutil.which("powershell.exe") or shutil.which("powershell")
    if not powershell:
        return False
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        target.unlink(missing_ok=True)
    except Exception:
        pass

    script = r"""$ErrorActionPreference = 'Stop'
$excel = $null
$book = $null
try {
    $source = [IO.Path]::GetFullPath($env:XVIEWER_XLS_SOURCE)
    $target = [IO.Path]::GetFullPath($env:XVIEWER_XLS_PDF)
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $excel.ScreenUpdating = $false
    try { $excel.AutomationSecurity = 3 } catch {}
    $book = $excel.Workbooks.Open($source, 0, $true)

    foreach ($ws in @($book.Worksheets)) {
        try {
            foreach ($shape in @($ws.Shapes)) {
                try { $shape.PrintObject = $true } catch {}
            }
        } catch {}
        try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($ws) } catch {}
    }

    # xlTypePDF=0, xlQualityStandard=0, IncludeDocProperties=True, IgnorePrintAreas=True
    $book.ExportAsFixedFormat(0, $target, 0, $true, $true)
}
finally {
    if ($book -ne $null) { try { $book.Close($false) } catch {} }
    if ($excel -ne $null) { try { $excel.Quit() } catch {} }
    if ($book -ne $null) { try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($book) } catch {} }
    if ($excel -ne $null) { try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($excel) } catch {} }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
"""
    env = os.environ.copy()
    env["XVIEWER_XLS_SOURCE"] = str(source)
    env["XVIEWER_XLS_PDF"] = str(target)
    try:
        completed = subprocess.run(
            _powershell_script_command(powershell, script),
            stdin=subprocess.DEVNULL,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=max(15, int(timeout)),
            env=env,
            **_hidden_subprocess_kwargs(),
        )
    except Exception:
        return False
    return completed.returncode == 0 and target.is_file() and target.stat().st_size > 100


def legacy_xls_pdf_preview_path(source: Path) -> Optional[Path]:
    """Return a cached Excel-rendered PDF for a legacy .xls workbook.

    The source stays read-only. This exact-view cache is used when old embedded
    pictures are not exposed by xlrd/openpyxl/legacy Shape extraction.
    """
    source = Path(source)
    if source.suffix.lower() != ".xls" or not source.is_file():
        return None
    target = _legacy_xls_pdf_path(source)
    if target.is_file() and target.stat().st_size > 100:
        return target
    with _LEGACY_XLS_CONVERSION_LOCK:
        if target.is_file() and target.stat().st_size > 100:
            return target
        if _render_xls_with_excel_pdf(source, target):
            return target
        try:
            target.unlink(missing_ok=True)
        except Exception:
            pass
    return None


def _hidden_subprocess_kwargs() -> dict[str, Any]:
    """Prevent helper conversions from flashing a console window on Windows."""
    if os.name != "nt":
        return {}
    startupinfo = subprocess.STARTUPINFO()
    startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
    startupinfo.wShowWindow = 0
    creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    return {"startupinfo": startupinfo, "creationflags": creationflags}


def _convert_xls_with_excel_com(source: Path, target: Path, timeout: int = 90) -> bool:
    """Use installed Microsoft Excel to create an image-capable .xlsx preview.

    The original .xls is never modified. Excel's own conversion is the most reliable
    way to preserve legacy BIFF/Escher drawing objects and embedded pictures, which
    xlrd does not expose to Python.
    """
    if os.name != "nt":
        return False
    powershell = shutil.which("powershell.exe") or shutil.which("powershell")
    if not powershell:
        return False

    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        target.unlink(missing_ok=True)
    except Exception:
        pass

    script = r'''$ErrorActionPreference = 'Stop'
$excel = $null
$book = $null
try {
    $source = [IO.Path]::GetFullPath($env:XVIEWER_XLS_SOURCE)
    $target = [IO.Path]::GetFullPath($env:XVIEWER_XLS_TARGET)
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    try { $excel.AutomationSecurity = 3 } catch {}
    $book = $excel.Workbooks.Open($source, 0, $true)
    $book.SaveAs($target, 51)
}
finally {
    if ($book -ne $null) { try { $book.Close($false) } catch {} }
    if ($excel -ne $null) { try { $excel.Quit() } catch {} }
    if ($book -ne $null) { try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($book) } catch {} }
    if ($excel -ne $null) { try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($excel) } catch {} }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
'''
    env = os.environ.copy()
    env["XVIEWER_XLS_SOURCE"] = str(source)
    env["XVIEWER_XLS_TARGET"] = str(target)
    try:
        completed = subprocess.run(
            _powershell_script_command(powershell, script),
            stdin=subprocess.DEVNULL,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=max(10, int(timeout)),
            env=env,
            **_hidden_subprocess_kwargs(),
        )
    except Exception:
        return False
    return completed.returncode == 0 and target.is_file() and target.stat().st_size > 0


def _extract_xls_images_with_excel_com(source: Path, manifest: Path, image_dir: Path, timeout: int = 120) -> bool:
    """Export picture-like Shapes from a legacy .xls workbook to PNG files.

    Some old BIFF workbooks keep photos as legacy Excel drawing shapes or metafiles.
    Even after Excel converts the workbook to .xlsx, openpyxl can still miss those
    objects.  Excel COM can render the original Shape itself, so xViewer stores a
    PNG sidecar plus worksheet/cell placement metadata and overlays it in the grid.
    """
    if os.name != "nt":
        return False
    powershell = shutil.which("powershell.exe") or shutil.which("powershell")
    if not powershell:
        return False

    manifest.parent.mkdir(parents=True, exist_ok=True)
    try:
        if image_dir.exists():
            shutil.rmtree(image_dir, ignore_errors=True)
        image_dir.mkdir(parents=True, exist_ok=True)
        manifest.unlink(missing_ok=True)
    except Exception:
        return False

    script = r'''$ErrorActionPreference = 'Stop'
$excel = $null
$book = $null
$records = New-Object System.Collections.ArrayList

function Export-ShapePng($ws, $shape, $dest) {
    $chartObj = $null
    try {
        $ws.Activate() | Out-Null
        try { $shape.Select($true) | Out-Null } catch {}
        try {
            $shape.CopyPicture(1, 2) | Out-Null
        } catch {
            try { $shape.Copy() | Out-Null } catch { return $false }
        }
        Start-Sleep -Milliseconds 80
        $w = [Math]::Max([double]$shape.Width, 2.0)
        $h = [Math]::Max([double]$shape.Height, 2.0)
        $chartObj = $ws.ChartObjects().Add(0, 0, $w, $h)
        $chart = $chartObj.Chart
        $chart.Paste() | Out-Null
        Start-Sleep -Milliseconds 80
        $ok = $chart.Export($dest, 'PNG')
        return [bool]$ok -and (Test-Path -LiteralPath $dest)
    } catch {
        return $false
    } finally {
        if ($chartObj -ne $null) {
            try { $chartObj.Delete() } catch {}
            try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($chartObj) } catch {}
        }
    }
}

try {
    $source = [IO.Path]::GetFullPath($env:XVIEWER_XLS_SOURCE)
    $manifest = [IO.Path]::GetFullPath($env:XVIEWER_XLS_MANIFEST)
    $outdir = [IO.Path]::GetFullPath($env:XVIEWER_XLS_IMAGE_DIR)
    [IO.Directory]::CreateDirectory($outdir) | Out-Null

    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $excel.ScreenUpdating = $false
    try { $excel.AutomationSecurity = 3 } catch {}
    $book = $excel.Workbooks.Open($source, 0, $true)

    $sheetIndex = 0
    foreach ($ws in @($book.Worksheets)) {
        $sheetIndex++
        $pictureIndex = 0
        foreach ($shape in @($ws.Shapes)) {
            try {
                $type = [int]$shape.Type
                $name = [string]$shape.Name
                # MsoShapeType: group=6, embedded OLE=7, linked OLE=10,
                # linked picture=11, picture=13. Some old product-photo workbooks
                # use the OLE variants instead of a normal picture shape.
                $pictureLike = $type -in @(6, 7, 10, 11, 13)
                if (-not $pictureLike -and $name -notmatch '^(Picture|Image|Photo)') { continue }
                if ([double]$shape.Width -lt 2 -or [double]$shape.Height -lt 2) { continue }

                $pictureIndex++
                $file = ('s{0:D3}_i{1:D4}.png' -f $sheetIndex, $pictureIndex)
                $dest = Join-Path $outdir $file
                if (-not (Export-ShapePng $ws $shape $dest)) { continue }

                $cell = $shape.TopLeftCell
                $row = [int]$cell.Row
                $column = [int]$cell.Column
                $xoff = [int][Math]::Round(([double]$shape.Left - [double]$cell.Left) * 96.0 / 72.0)
                $yoff = [int][Math]::Round(([double]$shape.Top - [double]$cell.Top) * 96.0 / 72.0)
                $width = [int][Math]::Max(1, [Math]::Round([double]$shape.Width * 96.0 / 72.0))
                $height = [int][Math]::Max(1, [Math]::Round([double]$shape.Height * 96.0 / 72.0))

                [void]$records.Add([pscustomobject]@{
                    sheet = [string]$ws.Name
                    row = $row
                    column = $column
                    x_offset = $xoff
                    y_offset = $yoff
                    width = $width
                    height = $height
                    file = $file
                    name = $name
                    shape_type = $type
                })
            } catch {
                continue
            }
        }
        try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($ws) } catch {}
    }

    $payload = [ordered]@{ version = 1; records = @($records) }
    $payload | ConvertTo-Json -Depth 6 | Set-Content -LiteralPath $manifest -Encoding UTF8
}
finally {
    if ($book -ne $null) { try { $book.Close($false) } catch {} }
    if ($excel -ne $null) { try { $excel.Quit() } catch {} }
    if ($book -ne $null) { try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($book) } catch {} }
    if ($excel -ne $null) { try { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($excel) } catch {} }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
'''
    env = os.environ.copy()
    env["XVIEWER_XLS_SOURCE"] = str(source)
    env["XVIEWER_XLS_MANIFEST"] = str(manifest)
    env["XVIEWER_XLS_IMAGE_DIR"] = str(image_dir)
    try:
        completed = subprocess.run(
            _powershell_script_command(powershell, script),
            stdin=subprocess.DEVNULL,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=max(15, int(timeout)),
            env=env,
            **_hidden_subprocess_kwargs(),
        )
    except Exception:
        return False
    return completed.returncode == 0 and manifest.is_file()


def _load_legacy_xls_image_manifest(source: Path) -> dict[str, list[EmbeddedImage]]:
    manifest = _legacy_xls_image_manifest_path(source)
    image_dir = _legacy_xls_image_dir(source)
    if not manifest.is_file():
        return {}
    try:
        payload = json.loads(manifest.read_text(encoding="utf-8-sig"))
        records = payload.get("records", []) if isinstance(payload, dict) else payload
        if isinstance(records, dict):
            records = [records]
    except Exception:
        return {}

    result: dict[str, list[EmbeddedImage]] = {}
    for number, record in enumerate(records or [], start=1):
        try:
            sheet = str(record.get("sheet", ""))
            filename = str(record.get("file", ""))
            image_path = image_dir / filename
            raw = image_path.read_bytes()
            if not sheet or not raw:
                continue
            item = EmbeddedImage(
                row=max(1, int(record.get("row", 1))),
                column=max(1, int(record.get("column", 1))),
                width=max(1, int(record.get("width", 1))),
                height=max(1, int(record.get("height", 1))),
                raw=raw,
                extension=image_path.suffix.lower() or ".png",
                label=str(record.get("name") or f"Legacy Image {number}"),
                x_offset=int(record.get("x_offset", 0) or 0),
                y_offset=int(record.get("y_offset", 0) or 0),
            )
            result.setdefault(sheet, []).append(item)
        except Exception:
            continue
    return result


def legacy_xls_images(source: Path) -> dict[str, list[EmbeddedImage]]:
    """Return picture shapes exported directly from a legacy .xls workbook."""
    source = Path(source)
    if source.suffix.lower() != ".xls" or not source.is_file():
        return {}
    manifest = _legacy_xls_image_manifest_path(source)
    if not manifest.is_file():
        with _LEGACY_XLS_CONVERSION_LOCK:
            if not manifest.is_file():
                _extract_xls_images_with_excel_com(source, manifest, _legacy_xls_image_dir(source))
    return _load_legacy_xls_image_manifest(source)


def _merge_embedded_images(primary: list[EmbeddedImage], extra: list[EmbeddedImage]) -> list[EmbeddedImage]:
    """Merge COM-exported legacy pictures without duplicating openpyxl pictures."""
    merged = list(primary)
    for candidate in extra:
        duplicate = False
        for current in merged:
            if (
                abs(current.row - candidate.row) <= 1
                and abs(current.column - candidate.column) <= 1
                and abs(current.width - candidate.width) <= max(8, int(candidate.width * 0.12))
                and abs(current.height - candidate.height) <= max(8, int(candidate.height * 0.12))
            ):
                duplicate = True
                break
        if not duplicate:
            merged.append(candidate)
    return merged


def _find_libreoffice() -> Optional[str]:
    for name in ("soffice", "soffice.exe", "libreoffice", "libreoffice.exe"):
        found = shutil.which(name)
        if found:
            return found
    if os.name == "nt":
        roots = [os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)")]
        for root in roots:
            if not root:
                continue
            for relative in (r"LibreOffice\program\soffice.exe", r"LibreOffice 7\program\soffice.exe"):
                candidate = Path(root) / relative
                if candidate.is_file():
                    return str(candidate)
    return None


def _convert_xls_with_libreoffice(source: Path, target: Path, timeout: int = 90) -> bool:
    """Fallback conversion when Microsoft Excel is unavailable."""
    soffice = _find_libreoffice()
    if not soffice:
        return False
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        with tempfile.TemporaryDirectory(prefix="xviewer-xls-", dir=str(target.parent)) as td:
            outdir = Path(td)
            completed = subprocess.run(
                [soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), str(source)],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=max(10, int(timeout)),
                **_hidden_subprocess_kwargs(),
            )
            if completed.returncode != 0:
                return False
            produced = outdir / f"{source.stem}.xlsx"
            if not produced.is_file():
                candidates = list(outdir.glob("*.xlsx"))
                if not candidates:
                    return False
                produced = candidates[0]
            shutil.move(str(produced), str(target))
            return target.is_file() and target.stat().st_size > 0
    except Exception:
        return False


def legacy_xls_preview_path(source: Path) -> Optional[Path]:
    """Return a cached .xlsx preview copy for a legacy .xls workbook.

    The preview is keyed by source path, size and modification time. The source .xls
    remains read-only. If neither Microsoft Excel nor LibreOffice can convert it,
    callers fall back to xlrd's cell-only legacy view.
    """
    source = Path(source)
    if source.suffix.lower() != ".xls" or not source.is_file():
        return None
    target = _legacy_xls_cache_path(source)
    if target.is_file() and target.stat().st_size > 0:
        # 2.6.3+ stores a second cache made from Excel-rendered legacy Shapes.
        # It is intentionally best-effort: the converted workbook remains usable
        # even when a particular old picture cannot be exported.
        if not _legacy_xls_image_manifest_path(source).is_file():
            with _LEGACY_XLS_CONVERSION_LOCK:
                if not _legacy_xls_image_manifest_path(source).is_file():
                    _extract_xls_images_with_excel_com(
                        source,
                        _legacy_xls_image_manifest_path(source),
                        _legacy_xls_image_dir(source),
                    )
        return target

    # Serialize conversions so fast LEFT-pane navigation does not launch several
    # Excel COM instances at once.
    with _LEGACY_XLS_CONVERSION_LOCK:
        if target.is_file() and target.stat().st_size > 0:
            return target
        target.parent.mkdir(parents=True, exist_ok=True)
        ok = _convert_xls_with_excel_com(source, target)
        if not ok:
            try:
                target.unlink(missing_ok=True)
            except Exception:
                pass
            ok = _convert_xls_with_libreoffice(source, target)
        if ok and target.is_file() and target.stat().st_size > 0:
            if not _legacy_xls_image_manifest_path(source).is_file():
                _extract_xls_images_with_excel_com(
                    source,
                    _legacy_xls_image_manifest_path(source),
                    _legacy_xls_image_dir(source),
                )
            return target
        try:
            target.unlink(missing_ok=True)
        except Exception:
            pass
        return None


def _column_name(index: int) -> str:
    result = ""
    number = max(1, int(index))
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _apply_tint(hex_color: str, tint: float) -> str:
    """Approximate Excel's tint transform for theme/indexed colors."""
    try:
        value = hex_color.lstrip("#")
        rgb = [int(value[i:i + 2], 16) for i in (0, 2, 4)]
        t = max(-1.0, min(1.0, float(tint or 0.0)))
        if t < 0:
            rgb = [round(c * (1.0 + t)) for c in rgb]
        elif t > 0:
            rgb = [round(c + (255 - c) * t) for c in rgb]
        return "#" + "".join(f"{max(0, min(255, c)):02X}" for c in rgb)
    except Exception:
        return hex_color


def _theme_palette(workbook: Any) -> list[str]:
    defaults = [
        "#000000", "#FFFFFF", "#1F497D", "#EEECE1",
        "#4F81BD", "#C0504D", "#9BBB59", "#8064A2",
        "#4BACC6", "#F79646", "#0000FF", "#800080",
    ]
    raw = getattr(workbook, "loaded_theme", None)
    if not raw:
        return defaults
    try:
        import xml.etree.ElementTree as ET
        root = ET.fromstring(raw)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = root.find(".//a:themeElements/a:clrScheme", ns)
        if scheme is None:
            return defaults
        colors: list[str] = []
        for child in list(scheme):
            color_node = next(iter(list(child)), None)
            if color_node is None:
                continue
            value = color_node.attrib.get("val") or color_node.attrib.get("lastClr")
            if value and len(value) >= 6:
                value = value[-6:]
                int(value, 16)
                colors.append(f"#{value.upper()}")
        if len(colors) >= 10:
            return (colors + defaults[len(colors):])[:12]
    except Exception:
        pass
    return defaults


def _rgb_from_openpyxl_color(color: Any, default: str, theme: Optional[list[str]] = None) -> str:
    try:
        if color is None:
            return default
        kind = getattr(color, "type", None)
        result = default
        if kind == "rgb":
            value = str(getattr(color, "rgb", "") or "")
            if len(value) == 8:
                value = value[2:]
            if len(value) == 6:
                int(value, 16)
                result = f"#{value.upper()}"
        elif kind == "theme" and theme:
            index = int(getattr(color, "theme", 0) or 0)
            if 0 <= index < len(theme):
                result = theme[index]
        elif kind == "indexed":
            from openpyxl.styles.colors import COLOR_INDEX
            index = int(getattr(color, "indexed", 0) or 0)
            if 0 <= index < len(COLOR_INDEX):
                value = str(COLOR_INDEX[index])
                if len(value) == 8:
                    value = value[2:]
                if len(value) == 6:
                    int(value, 16)
                    result = f"#{value.upper()}"
        tint = float(getattr(color, "tint", 0.0) or 0.0)
        return _apply_tint(result, tint) if tint else result
    except Exception:
        return default


def _excel_column_width_to_pixels(width: float, zoom: float = 1.0) -> int:
    """Close approximation of Excel's character-width to pixel conversion."""
    try:
        value = max(0.0, float(width))
        pixels = int(value * 12 + 0.5) if value < 1 else int(value * 7 + 5)
        return max(0, int(round(pixels * zoom)))
    except Exception:
        return max(0, int(round(DEFAULT_COLUMN_WIDTH * zoom)))


def _emu_to_pixels(value: Any) -> int:
    try:
        return int(round(float(value or 0) / 9525.0))
    except Exception:
        return 0


def _format_number(value: float | int, number_format: str) -> str:
    fmt = (number_format or "").strip()
    if not fmt or fmt.lower() == "general":
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)

    lower = fmt.lower()
    if "%" in fmt:
        decimals = 0
        if "." in fmt:
            decimals = sum(1 for c in fmt.split(".", 1)[1].split("%", 1)[0] if c in "0#")
        return f"{float(value) * 100:.{decimals}f}%"

    decimals = 0
    if "." in fmt:
        tail = fmt.split(".", 1)[1]
        decimals = sum(1 for c in tail if c in "0#")
        decimals = min(decimals, 8)

    use_group = "," in fmt.split(".", 1)[0]
    symbol = ""
    for candidate in ("$", "€", "£", "¥", "₩"):
        if candidate in fmt:
            symbol = candidate
            break

    try:
        if decimals:
            number = f"{float(value):,.{decimals}f}" if use_group else f"{float(value):.{decimals}f}"
        elif use_group:
            number = f"{float(value):,.0f}"
        else:
            number = str(int(value)) if float(value).is_integer() else str(value)
        return f"{symbol}{number}"
    except Exception:
        return str(value)


def _format_excel_date_like(value: Any, number_format: str) -> str:
    fmt = (number_format or "").lower().strip()
    # Remove the most common Excel decorations while keeping separators.
    import re
    fmt = re.sub(r"\[\$-[^\]]+\]", "", fmt)
    fmt = fmt.replace("\\-", "-").replace("\\/", "/")
    fmt = re.sub(r'"([^"]*)"', r"\1", fmt)
    date_patterns = [
        ("mm-dd-yyyy", "%m-%d-%Y"), ("m-d-yyyy", "%m-%d-%Y"),
        ("mm-dd-yy", "%m-%d-%y"), ("m-d-yy", "%m-%d-%y"),
        ("mm/dd/yyyy", "%m/%d/%Y"), ("m/d/yyyy", "%m/%d/%Y"),
        ("mm/dd/yy", "%m/%d/%y"), ("m/d/yy", "%m/%d/%y"),
        ("yyyy-mm-dd", "%Y-%m-%d"), ("yyyy/m/d", "%Y/%m/%d"),
        ("dd-mmm-yy", "%d-%b-%y"), ("d-mmm-yy", "%d-%b-%y"),
        ("mmm d, yyyy", "%b %d, %Y"), ("mmmm d, yyyy", "%B %d, %Y"),
    ]
    time_patterns = [
        ("h:mm:ss am/pm", "%I:%M:%S %p"), ("hh:mm:ss am/pm", "%I:%M:%S %p"),
        ("h:mm am/pm", "%I:%M %p"), ("hh:mm am/pm", "%I:%M %p"),
        ("hh:mm:ss", "%H:%M:%S"), ("h:mm:ss", "%H:%M:%S"),
        ("hh:mm", "%H:%M"), ("h:mm", "%H:%M"),
    ]
    for excel_fmt, py_fmt in date_patterns:
        if excel_fmt in fmt:
            return value.strftime(py_fmt)
    for excel_fmt, py_fmt in time_patterns:
        if excel_fmt in fmt:
            return value.strftime(py_fmt).lstrip("0") if "%I" in py_fmt else value.strftime(py_fmt)
    if isinstance(value, datetime):
        if "h" in fmt or "s" in fmt:
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return str(value)


def format_cell_value(value: Any, number_format: str = "") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return _format_excel_date_like(value, number_format)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return _format_number(value, number_format)
    return str(value)


@dataclass(frozen=True)
class BorderSide:
    style: str = ""
    color: str = "#000000"


@dataclass(frozen=True)
class CellStyle:
    fill: str = "#ffffff"
    foreground: str = "#202020"
    bold: bool = False
    italic: bool = False
    underline: bool = False
    font_name: str = "Calibri"
    font_size: int = 11
    horizontal: str = "general"
    vertical: str = "bottom"
    wrap: bool = False
    indent: int = 0
    left: BorderSide = BorderSide()
    right: BorderSide = BorderSide()
    top: BorderSide = BorderSide()
    bottom: BorderSide = BorderSide()


@dataclass(frozen=True)
class EmbeddedImage:
    row: int
    column: int
    width: int
    height: int
    raw: bytes
    extension: str = ".png"
    label: str = "Image"
    x_offset: int = 0
    y_offset: int = 0
    end_row: int | None = None
    end_column: int | None = None
    end_x_offset: int = 0
    end_y_offset: int = 0


class AxisMetrics:
    """Sparse row/column sizes without allocating an entry for every index."""

    def __init__(self, count: int, default_size: int, overrides: dict[int, int]) -> None:
        self.count = max(1, int(count))
        self.default_size = max(1, int(default_size))
        cleaned = {
            int(index): max(0, int(size))
            for index, size in overrides.items()
            if 1 <= int(index) <= self.count and int(size) != self.default_size
        }
        self.overrides = cleaned
        self._indices = sorted(cleaned)
        self._prefix: list[int] = [0]
        total = 0
        for index in self._indices:
            total += cleaned[index] - self.default_size
            self._prefix.append(total)

    def _delta_before(self, index: int) -> int:
        pos = bisect.bisect_left(self._indices, index)
        return self._prefix[pos]

    def start(self, index: int) -> int:
        index = max(1, min(self.count + 1, int(index)))
        return (index - 1) * self.default_size + self._delta_before(index)

    def size(self, index: int) -> int:
        return self.overrides.get(int(index), self.default_size)

    @property
    def total(self) -> int:
        return self.start(self.count + 1)

    def index_at(self, coordinate: float) -> int:
        if coordinate <= 0:
            candidate = 1
        elif coordinate >= self.total:
            candidate = self.count
        else:
            lo, hi = 1, self.count
            candidate = 1
            while lo <= hi:
                mid = (lo + hi) // 2
                start = self.start(mid)
                end = start + self.size(mid)
                if coordinate < start:
                    hi = mid - 1
                elif coordinate >= end:
                    lo = mid + 1
                    candidate = min(self.count, lo)
                else:
                    candidate = mid
                    break
        if self.size(candidate) > 0:
            return candidate
        # Hidden rows/columns occupy no pixels. Prefer the next visible item,
        # then fall back to the previous visible item.
        for idx in range(candidate + 1, self.count + 1):
            if self.size(idx) > 0:
                return idx
        for idx in range(candidate - 1, 0, -1):
            if self.size(idx) > 0:
                return idx
        return max(1, min(self.count, candidate))

    def next_visible(self, index: int, direction: int) -> int:
        step = 1 if direction >= 0 else -1
        current = max(1, min(self.count, int(index)))
        candidate = current + step
        while 1 <= candidate <= self.count:
            if self.size(candidate) > 0:
                return candidate
            candidate += step
        return current

    def first_visible(self) -> int:
        for index in range(1, self.count + 1):
            if self.size(index) > 0:
                return index
        return 1

    def last_visible(self) -> int:
        for index in range(self.count, 0, -1):
            if self.size(index) > 0:
                return index
        return self.count


class SheetAdapter:
    title: str
    max_row: int
    max_column: int
    merged_ranges: list[tuple[int, int, int, int]]
    images: list[EmbeddedImage]

    def value(self, row: int, column: int) -> Any:
        raise NotImplementedError

    def text(self, row: int, column: int) -> str:
        raise NotImplementedError

    def style(self, row: int, column: int) -> CellStyle:
        return CellStyle()

    def row_overrides(self, zoom: float = 1.0) -> dict[int, int]:
        return {}

    def column_overrides(self, zoom: float = 1.0) -> dict[int, int]:
        return {}

    def default_row_size(self, zoom: float = 1.0) -> int:
        return max(1, int(round(DEFAULT_ROW_HEIGHT * zoom)))

    def default_column_size(self, zoom: float = 1.0) -> int:
        return max(1, int(round(DEFAULT_COLUMN_WIDTH * zoom)))

    @property
    def show_gridlines(self) -> bool:
        return True


class OpenPyxlSheetAdapter(SheetAdapter):
    def __init__(self, worksheet, images: list[EmbeddedImage]) -> None:
        self.ws = worksheet
        self.title = worksheet.title
        self._theme = _theme_palette(worksheet.parent)
        self._style_cache: dict[int, CellStyle] = {}
        self.max_row = max(1, int(worksheet.max_row or 1))
        self.max_column = max(1, int(worksheet.max_column or 1))
        for image in images:
            self.max_row = max(self.max_row, image.end_row or image.row)
            self.max_column = max(self.max_column, image.end_column or image.column)
        self.images = images
        self.merged_ranges = [
            (item.min_col, item.min_row, item.max_col, item.max_row)
            for item in worksheet.merged_cells.ranges
        ]

    @property
    def show_gridlines(self) -> bool:
        try:
            return getattr(self.ws.sheet_view, "showGridLines", None) is not False
        except Exception:
            return True

    def default_row_size(self, zoom: float = 1.0) -> int:
        try:
            points = float(getattr(self.ws.sheet_format, "defaultRowHeight", None) or 15.0)
            pixels = int(round(points * 96.0 / 72.0 * zoom))
            return max(1, min(MAX_ROW_HEIGHT, pixels))
        except Exception:
            return super().default_row_size(zoom)

    def default_column_size(self, zoom: float = 1.0) -> int:
        try:
            width = getattr(self.ws.sheet_format, "defaultColWidth", None)
            if width is None:
                width = 8.43
            return max(1, min(MAX_COLUMN_WIDTH, _excel_column_width_to_pixels(float(width), zoom)))
        except Exception:
            return super().default_column_size(zoom)

    def value(self, row: int, column: int) -> Any:
        return self.ws.cell(row, column).value

    def text(self, row: int, column: int) -> str:
        cell = self.ws.cell(row, column)
        return format_cell_value(cell.value, getattr(cell, "number_format", "") or "")

    def _border_side(self, side: Any) -> BorderSide:
        style = str(getattr(side, "style", "") or "")
        color = _rgb_from_openpyxl_color(getattr(side, "color", None), "#000000", self._theme)
        return BorderSide(style=style, color=color)

    def style(self, row: int, column: int) -> CellStyle:
        cell = self.ws.cell(row, column)
        style_id = int(getattr(cell, "style_id", 0) or 0)
        cached = self._style_cache.get(style_id)
        if cached is not None:
            return cached

        fill = "#ffffff"
        try:
            if getattr(cell.fill, "fill_type", None):
                fill = _rgb_from_openpyxl_color(cell.fill.fgColor, fill, self._theme)
        except Exception:
            pass
        foreground = _rgb_from_openpyxl_color(getattr(cell.font, "color", None), "#202020", self._theme)

        horizontal = (getattr(cell.alignment, "horizontal", None) or "general").lower()
        if horizontal in {"distributed", "fill", "justify", "centercontinuous", "centercontinuous"}:
            horizontal = "center" if horizontal.startswith("center") else "left"
        if horizontal not in {"left", "center", "right", "general"}:
            horizontal = "general"
        vertical = (getattr(cell.alignment, "vertical", None) or "bottom").lower()
        if vertical not in {"top", "center", "bottom"}:
            vertical = "bottom"

        size = getattr(cell.font, "sz", None) or 11
        try:
            size_i = max(6, min(72, int(round(float(size)))))
        except Exception:
            size_i = 11
        font_name = str(getattr(cell.font, "name", None) or "Calibri")
        underline = bool(getattr(cell.font, "u", None) not in {None, False, "none"})
        indent = int(getattr(cell.alignment, "indent", 0) or 0)

        border = getattr(cell, "border", None)
        result = CellStyle(
            fill=fill,
            foreground=foreground,
            bold=bool(getattr(cell.font, "bold", False)),
            italic=bool(getattr(cell.font, "italic", False)),
            underline=underline,
            font_name=font_name,
            font_size=size_i,
            horizontal=horizontal,
            vertical=vertical,
            wrap=bool(getattr(cell.alignment, "wrap_text", False)),
            indent=max(0, min(15, indent)),
            left=self._border_side(getattr(border, "left", None)),
            right=self._border_side(getattr(border, "right", None)),
            top=self._border_side(getattr(border, "top", None)),
            bottom=self._border_side(getattr(border, "bottom", None)),
        )
        self._style_cache[style_id] = result
        return result

    def row_overrides(self, zoom: float = 1.0) -> dict[int, int]:
        out: dict[int, int] = {}
        for index, dimension in self.ws.row_dimensions.items():
            try:
                idx = int(index)
            except Exception:
                continue
            if bool(getattr(dimension, "hidden", False)):
                out[idx] = 0
                continue
            height = getattr(dimension, "height", None)
            if height is None:
                continue
            try:
                pixels = int(round(float(height) * 96.0 / 72.0 * zoom))
                out[idx] = max(1, min(MAX_ROW_HEIGHT, pixels))
            except Exception:
                pass
        return out

    def column_overrides(self, zoom: float = 1.0) -> dict[int, int]:
        from openpyxl.utils.cell import column_index_from_string

        out: dict[int, int] = {}
        for name, dimension in self.ws.column_dimensions.items():
            try:
                start = int(getattr(dimension, "min", 0) or 0)
                stop = int(getattr(dimension, "max", 0) or 0)
                if start <= 0:
                    start = column_index_from_string(str(name))
                if stop < start:
                    stop = start
            except Exception:
                continue
            if bool(getattr(dimension, "hidden", False)):
                for index in range(start, min(self.max_column, stop) + 1):
                    out[index] = 0
                continue
            width = getattr(dimension, "width", None)
            if width is None:
                continue
            try:
                pixels = _excel_column_width_to_pixels(float(width), zoom)
                pixels = max(1, min(MAX_COLUMN_WIDTH, pixels))
                for index in range(start, min(self.max_column, stop) + 1):
                    out[index] = pixels
            except Exception:
                pass
        return out


class XlrdSheetAdapter(SheetAdapter):
    def __init__(self, sheet, images: Optional[list[EmbeddedImage]] = None) -> None:
        self.sheet = sheet
        self.title = sheet.name
        self.images = list(images or [])
        self.max_row = max(1, int(sheet.nrows or 1))
        self.max_column = max(1, int(sheet.ncols or 1))
        for image in self.images:
            self.max_row = max(self.max_row, image.end_row or image.row)
            self.max_column = max(self.max_column, image.end_column or image.column)
        self.merged_ranges = [
            (clo + 1, rlo + 1, chi, rhi)
            for rlo, rhi, clo, chi in getattr(sheet, "merged_cells", [])
        ]

    def value(self, row: int, column: int) -> Any:
        if row > self.sheet.nrows or column > self.sheet.ncols:
            return None
        return self.sheet.cell_value(row - 1, column - 1)

    def text(self, row: int, column: int) -> str:
        value = self.value(row, column)
        if value == "":
            return ""
        return format_cell_value(value)


class WorkbookModel:
    def __init__(self, path: Path, show_formulas: bool = False) -> None:
        self.path = path
        self.show_formulas = show_formulas
        self.workbook = None
        self.sheets: list[SheetAdapter] = []
        self.kind = "Excel"
        self.skipped_image_errors: list[str] = []
        self._load()

    @staticmethod
    def _modern_images(worksheet) -> list[EmbeddedImage]:
        results: list[EmbeddedImage] = []
        for number, image in enumerate(getattr(worksheet, "_images", []), start=1):
            try:
                anchor = getattr(image, "anchor", None)
                marker = getattr(anchor, "_from", None)
                if marker is None:
                    continue
                row = int(marker.row) + 1
                column = int(marker.col) + 1
                x_offset = _emu_to_pixels(getattr(marker, "colOff", 0))
                y_offset = _emu_to_pixels(getattr(marker, "rowOff", 0))
                raw = image._data()
                if not raw:
                    continue
                extension = ".png"
                try:
                    extension = "." + str(getattr(image, "format", "png") or "png").lower().lstrip(".")
                except Exception:
                    pass
                width = max(1, int(round(float(getattr(image, "width", 96) or 96))))
                height = max(1, int(round(float(getattr(image, "height", 96) or 96))))
                ext = getattr(anchor, "ext", None)
                if ext is not None:
                    width = max(1, _emu_to_pixels(getattr(ext, "cx", 0)) or width)
                    height = max(1, _emu_to_pixels(getattr(ext, "cy", 0)) or height)
                to_marker = getattr(anchor, "_to", None)
                end_row = end_column = None
                end_x_offset = end_y_offset = 0
                if to_marker is not None:
                    end_row = int(to_marker.row) + 1
                    end_column = int(to_marker.col) + 1
                    end_x_offset = _emu_to_pixels(getattr(to_marker, "colOff", 0))
                    end_y_offset = _emu_to_pixels(getattr(to_marker, "rowOff", 0))
                results.append(
                    EmbeddedImage(
                        row=row,
                        column=column,
                        width=width,
                        height=height,
                        raw=bytes(raw),
                        extension=extension,
                        label=f"Image {number}",
                        x_offset=x_offset,
                        y_offset=y_offset,
                        end_row=end_row,
                        end_column=end_column,
                        end_x_offset=end_x_offset,
                        end_y_offset=end_y_offset,
                    )
                )
            except Exception:
                continue
        return results

    def _load(self) -> None:
        suffix = self.path.suffix.lower()
        if suffix in MODERN_EXCEL_EXTENSIONS:
            keep_vba = suffix in {".xlsm", ".xltm"}
            workbook, skipped_images = _safe_openpyxl_load_workbook(
                self.path,
                data_only=not self.show_formulas,
                read_only=False,
                keep_vba=keep_vba,
            )
            self.skipped_image_errors = skipped_images
            self.workbook = workbook
            self.sheets = [
                OpenPyxlSheetAdapter(ws, self._modern_images(ws))
                for ws in workbook.worksheets
            ]
            self.kind = "Modern Excel"
            return

        if suffix == ".xls":
            preview = legacy_xls_preview_path(self.path)
            legacy_by_sheet = legacy_xls_images(self.path)
            if preview is not None:
                workbook, skipped_images = _safe_openpyxl_load_workbook(
                    preview,
                    data_only=not self.show_formulas,
                    read_only=False,
                    keep_vba=False,
                )
                self.skipped_image_errors = skipped_images
                self.workbook = workbook
                self.sheets = [
                    OpenPyxlSheetAdapter(
                        ws,
                        _merge_embedded_images(self._modern_images(ws), legacy_by_sheet.get(ws.title, [])),
                    )
                    for ws in workbook.worksheets
                ]
                self.kind = "Legacy Excel (.xls) - converted preview + direct picture export"
                return

            import xlrd

            workbook = xlrd.open_workbook(self.path, on_demand=True, formatting_info=False)
            self.workbook = workbook
            self.sheets = [
                XlrdSheetAdapter(
                    workbook.sheet_by_index(i),
                    legacy_by_sheet.get(workbook.sheet_by_index(i).name, []),
                )
                for i in range(workbook.nsheets)
            ]
            self.kind = "Legacy Excel (.xls) - cell fallback + direct picture export"
            return

        raise ValueError(f"Unsupported Excel format: {suffix}")

    def close(self) -> None:
        if self.workbook is None:
            return
        try:
            close = getattr(self.workbook, "close", None)
            if callable(close):
                close()
        except Exception:
            pass
        try:
            release = getattr(self.workbook, "release_resources", None)
            if callable(release):
                release()
        except Exception:
            pass
        self.workbook = None


def copy_pil_image_to_windows_clipboard(image: Image.Image) -> None:
    if os.name != "nt":
        raise RuntimeError("Image clipboard is supported on Windows in this build.")

    import ctypes

    CF_DIB = 8
    GMEM_MOVEABLE = 0x0002
    kernel32 = ctypes.windll.kernel32
    user32 = ctypes.windll.user32
    kernel32.GlobalAlloc.argtypes = [ctypes.c_uint, ctypes.c_size_t]
    kernel32.GlobalAlloc.restype = ctypes.c_void_p
    kernel32.GlobalLock.argtypes = [ctypes.c_void_p]
    kernel32.GlobalLock.restype = ctypes.c_void_p
    kernel32.GlobalUnlock.argtypes = [ctypes.c_void_p]
    kernel32.GlobalUnlock.restype = ctypes.c_int
    kernel32.GlobalFree.argtypes = [ctypes.c_void_p]
    kernel32.GlobalFree.restype = ctypes.c_void_p
    user32.OpenClipboard.argtypes = [ctypes.c_void_p]
    user32.OpenClipboard.restype = ctypes.c_int
    user32.EmptyClipboard.restype = ctypes.c_int
    user32.SetClipboardData.argtypes = [ctypes.c_uint, ctypes.c_void_p]
    user32.SetClipboardData.restype = ctypes.c_void_p
    user32.CloseClipboard.restype = ctypes.c_int

    converted = image.convert("RGB")
    stream = io.BytesIO()
    converted.save(stream, "BMP")
    data = stream.getvalue()[14:]

    if not user32.OpenClipboard(None):
        raise RuntimeError("Could not open the Windows clipboard.")
    handle = None
    try:
        if not user32.EmptyClipboard():
            raise RuntimeError("Could not clear the Windows clipboard.")
        handle = kernel32.GlobalAlloc(GMEM_MOVEABLE, len(data))
        if not handle:
            raise RuntimeError("Could not allocate clipboard memory.")
        pointer = kernel32.GlobalLock(handle)
        if not pointer:
            raise RuntimeError("Could not lock clipboard memory.")
        try:
            ctypes.memmove(pointer, data, len(data))
        finally:
            kernel32.GlobalUnlock(handle)
        if not user32.SetClipboardData(CF_DIB, handle):
            raise RuntimeError("Could not place the image on the clipboard.")
        handle = None
    finally:
        user32.CloseClipboard()
        if handle:
            kernel32.GlobalFree(handle)


def _sheet_content_extent(sheet: SheetAdapter, rows: AxisMetrics, columns: AxisMetrics, zoom: float) -> tuple[int, int]:
    """Return scrollable pixel bounds, including images that extend beyond used cells."""
    max_x = int(columns.total)
    max_y = int(rows.total)
    margin = max(24, int(round(24 * float(zoom))))
    for record in sheet.images:
        x = columns.start(record.column) + int(round(record.x_offset * zoom))
        y = rows.start(record.row) + int(round(record.y_offset * zoom))
        if record.end_column is not None and record.end_row is not None:
            x2 = columns.start(record.end_column) + int(round(record.end_x_offset * zoom))
            y2 = rows.start(record.end_row) + int(round(record.end_y_offset * zoom))
        else:
            x2 = x + max(1, int(round(record.width * zoom)))
            y2 = y + max(1, int(round(record.height * zoom)))
        max_x = max(max_x, x2 + margin)
        max_y = max(max_y, y2 + margin)
    return max(1, max_x), max(1, max_y)


class VirtualSheet(ttk.Frame):
    def __init__(self, master, sheet: SheetAdapter, status_callback, ui_palette: Optional[Mapping[str, str]] = None) -> None:
        super().__init__(master)
        self.sheet = sheet
        self.ui_palette = dict(ui_palette or LIGHT_PALETTE)
        self.status_callback = status_callback
        self.zoom = 1.0
        self.active_cell = (1, 1)
        self.anchor_cell = (1, 1)
        self.selected_image: Optional[int] = None
        self._font_cache: dict[tuple[str, int, bool, bool, bool], tkfont.Font] = {}
        self._visible_images: dict[int, ImageTk.PhotoImage] = {}
        self._render_pending: Optional[str] = None
        self._build_metrics()
        self.active_cell = (self.rows.first_visible(), self.columns.first_visible())
        self.anchor_cell = self.active_cell
        self._build_ui()
        self.after_idle(self.render)

    def _build_metrics(self) -> None:
        default_row = max(1, int(self.sheet.default_row_size(self.zoom)))
        default_col = max(1, int(self.sheet.default_column_size(self.zoom)))
        self.rows = AxisMetrics(self.sheet.max_row, default_row, self.sheet.row_overrides(self.zoom))
        self.columns = AxisMetrics(self.sheet.max_column, default_col, self.sheet.column_overrides(self.zoom))
        self.content_width, self.content_height = _sheet_content_extent(
            self.sheet, self.rows, self.columns, self.zoom
        )

    def _build_ui(self) -> None:
        self.rowconfigure(1, weight=1)
        self.columnconfigure(1, weight=1)

        header_bg = self.ui_palette.get("header", HEADER_BACKGROUND)
        canvas_bg = self.ui_palette.get("canvas_margin", "white")
        self.corner = tk.Canvas(self, width=52, height=28, highlightthickness=0, background=header_bg)
        self.col_header = tk.Canvas(self, height=28, highlightthickness=0, background=header_bg)
        self.row_header = tk.Canvas(self, width=52, highlightthickness=0, background=header_bg)
        self.body = tk.Canvas(self, highlightthickness=0, background=canvas_bg, takefocus=True)
        self.xbar = ttk.Scrollbar(self, orient="horizontal", command=self._xview)
        self.ybar = ttk.Scrollbar(self, orient="vertical", command=self._yview)

        self.corner.grid(row=0, column=0, sticky="nsew")
        self.col_header.grid(row=0, column=1, sticky="ew")
        self.row_header.grid(row=1, column=0, sticky="ns")
        self.body.grid(row=1, column=1, sticky="nsew")
        self.ybar.grid(row=1, column=2, sticky="ns")
        self.xbar.grid(row=2, column=1, sticky="ew")

        self.body.configure(
            scrollregion=(0, 0, self.content_width, self.content_height),
            xscrollcommand=self._xscroll_set,
            yscrollcommand=self._yscroll_set,
        )
        self.col_header.configure(scrollregion=(0, 0, self.content_width, 28))
        self.row_header.configure(scrollregion=(0, 0, 52, self.content_height))

        self.corner.create_rectangle(0, 0, 52, 28, fill=header_bg, outline=self.ui_palette.get("header_border", "#bdbdbd"))

        self.body.bind("<Configure>", lambda _event: self.schedule_render())
        self.body.bind("<Button-1>", self._mouse_down)
        self.body.bind("<B1-Motion>", self._mouse_drag)
        self.body.bind("<ButtonRelease-1>", self._mouse_up)
        self.body.bind("<Button-3>", self._right_click)
        self.body.bind("<MouseWheel>", self._mouse_wheel)
        self.body.bind("<Shift-MouseWheel>", self._shift_mouse_wheel)
        self.body.bind("<Control-c>", self.copy_selection)
        self.body.bind("<Control-C>", self.copy_selection)
        self.body.bind("<KeyPress>", self._key_press)


    def apply_ui_theme(self, palette: Mapping[str, str]) -> None:
        """Retheme viewer chrome while preserving workbook cell formatting."""
        self.ui_palette = dict(palette)
        header_bg = self.ui_palette.get("header", HEADER_BACKGROUND)
        canvas_bg = self.ui_palette.get("canvas_margin", "white")
        for canvas in (self.corner, self.col_header, self.row_header):
            try:
                canvas.configure(background=header_bg)
            except Exception:
                pass
        try:
            self.body.configure(background=canvas_bg)
        except Exception:
            pass
        try:
            self.corner.delete("all")
            self.corner.create_rectangle(
                0, 0, 52, 28,
                fill=header_bg,
                outline=self.ui_palette.get("header_border", "#bdbdbd"),
            )
        except Exception:
            pass
        self.schedule_render()

    def _xscroll_set(self, first: str, last: str) -> None:
        self.xbar.set(first, last)
        self.col_header.xview_moveto(float(first))
        self.schedule_render()

    def _yscroll_set(self, first: str, last: str) -> None:
        self.ybar.set(first, last)
        self.row_header.yview_moveto(float(first))
        self.schedule_render()

    def _xview(self, *args) -> None:
        self.body.xview(*args)
        self.col_header.xview(*args)
        self.schedule_render()

    def _yview(self, *args) -> None:
        self.body.yview(*args)
        self.row_header.yview(*args)
        self.schedule_render()

    def _mouse_wheel(self, event) -> str:
        amount = -1 if event.delta > 0 else 1
        self.body.yview_scroll(amount * 3, "units")
        self.schedule_render()
        return "break"

    def _shift_mouse_wheel(self, event) -> str:
        amount = -1 if event.delta > 0 else 1
        self.body.xview_scroll(amount * 3, "units")
        self.schedule_render()
        return "break"

    def schedule_render(self) -> None:
        if self._render_pending is not None:
            return
        self._render_pending = self.after_idle(self._render_scheduled)

    def _render_scheduled(self) -> None:
        self._render_pending = None
        self.render()

    def _visible_bounds(self) -> tuple[int, int, int, int, float, float, float, float]:
        left = self.body.canvasx(0)
        top = self.body.canvasy(0)
        width = max(1, self.body.winfo_width())
        height = max(1, self.body.winfo_height())
        right = left + width
        bottom = top + height
        first_col = self.columns.index_at(left)
        last_col = self.columns.index_at(right)
        first_row = self.rows.index_at(top)
        last_row = self.rows.index_at(bottom)
        first_col = max(1, first_col - 1)
        first_row = max(1, first_row - 1)
        last_col = min(self.sheet.max_column, last_col + 1)
        last_row = min(self.sheet.max_row, last_row + 1)
        return first_row, last_row, first_col, last_col, left, top, right, bottom

    def _font(self, style: CellStyle) -> tkfont.Font:
        size = max(6, int(round(style.font_size * self.zoom)))
        family = style.font_name or "Calibri"
        key = (family, size, style.bold, style.italic, style.underline)
        cached = self._font_cache.get(key)
        if cached is not None:
            return cached
        weight = "bold" if style.bold else "normal"
        slant = "italic" if style.italic else "roman"
        font = tkfont.Font(
            family=family,
            size=size,
            weight=weight,
            slant=slant,
            underline=1 if style.underline else 0,
        )
        self._font_cache[key] = font
        return font

    @staticmethod
    def _border_width(style: str) -> int:
        value = (style or "").lower()
        if not value:
            return 0
        if value in {"medium", "mediumdashed", "mediumdashdot", "mediumdashdotdot"}:
            return 2
        if value in {"thick", "double"}:
            return 3
        return 1

    def _draw_border_side(self, x0: float, y0: float, x1: float, y1: float, side: BorderSide, which: str) -> None:
        width = self._border_width(side.style)
        if width <= 0:
            return
        if which == "left":
            coords = (x0, y0, x0, y1)
        elif which == "right":
            coords = (x1, y0, x1, y1)
        elif which == "top":
            coords = (x0, y0, x1, y0)
        else:
            coords = (x0, y1, x1, y1)
        self.body.create_line(*coords, fill=side.color or "#000000", width=width)
        if (side.style or "").lower() == "double":
            offset = 2
            if which == "left":
                coords = (x0 + offset, y0, x0 + offset, y1)
            elif which == "right":
                coords = (x1 - offset, y0, x1 - offset, y1)
            elif which == "top":
                coords = (x0, y0 + offset, x1, y0 + offset)
            else:
                coords = (x0, y1 - offset, x1, y1 - offset)
            self.body.create_line(*coords, fill=side.color or "#000000", width=1)

    def _merged_visible(
        self,
        first_row: int,
        last_row: int,
        first_col: int,
        last_col: int,
    ) -> list[tuple[int, int, int, int]]:
        out = []
        for min_col, min_row, max_col, max_row in self.sheet.merged_ranges:
            if max_row < first_row or min_row > last_row or max_col < first_col or min_col > last_col:
                continue
            out.append((min_col, min_row, max_col, max_row))
        return out

    def _cell_anchor(self, row: int, column: int) -> tuple[int, int, int, int] | None:
        for min_col, min_row, max_col, max_row in self.sheet.merged_ranges:
            if min_row <= row <= max_row and min_col <= column <= max_col:
                return (min_row, min_col, max_row, max_col)
        return None

    def render(self) -> None:
        if not self.body.winfo_exists():
            return
        self.body.delete("all")
        self.col_header.delete("all")
        self.row_header.delete("all")
        self._visible_images.clear()

        first_row, last_row, first_col, last_col, left, top, right, bottom = self._visible_bounds()
        merged = self._merged_visible(first_row, last_row, first_col, last_col)
        covered: set[tuple[int, int]] = set()
        for min_col, min_row, max_col, max_row in merged:
            for row in range(max(first_row, min_row), min(last_row, max_row) + 1):
                for col in range(max(first_col, min_col), min(last_col, max_col) + 1):
                    covered.add((row, col))

        header_font = self._font(CellStyle(font_name="Segoe UI", bold=True, font_size=9))
        for col in range(first_col, last_col + 1):
            width = self.columns.size(col)
            if width <= 0:
                continue
            x0 = self.columns.start(col)
            x1 = x0 + width
            self.col_header.create_rectangle(x0, 0, x1, 28, fill=self.ui_palette.get("header", HEADER_BACKGROUND), outline=self.ui_palette.get("header_border", "#bdbdbd"))
            self.col_header.create_text((x0 + x1) / 2, 14, text=_column_name(col), fill=self.ui_palette.get("header_foreground", HEADER_FOREGROUND), font=header_font)

        for row in range(first_row, last_row + 1):
            height = self.rows.size(row)
            if height <= 0:
                continue
            y0 = self.rows.start(row)
            y1 = y0 + height
            self.row_header.create_rectangle(0, y0, 52, y1, fill=self.ui_palette.get("header", HEADER_BACKGROUND), outline=self.ui_palette.get("header_border", "#bdbdbd"))
            self.row_header.create_text(47, (y0 + y1) / 2, text=str(row), anchor="e", fill=self.ui_palette.get("header_foreground", HEADER_FOREGROUND), font=header_font)

        cells: list[tuple[int, int, float, float, float, float]] = []
        for row in range(first_row, last_row + 1):
            if self.rows.size(row) <= 0:
                continue
            for col in range(first_col, last_col + 1):
                if self.columns.size(col) <= 0 or (row, col) in covered:
                    continue
                x0 = self.columns.start(col)
                y0 = self.rows.start(row)
                x1 = x0 + self.columns.size(col)
                y1 = y0 + self.rows.size(row)
                cells.append((row, col, x0, y0, x1, y1))

        for min_col, min_row, max_col, max_row in merged:
            # A merged region whose visible width/height is zero should not render.
            x0 = self.columns.start(min_col)
            y0 = self.rows.start(min_row)
            x1 = self.columns.start(max_col) + self.columns.size(max_col)
            y1 = self.rows.start(max_row) + self.rows.size(max_row)
            if x1 > x0 and y1 > y0:
                cells.append((min_row, min_col, x0, y0, x1, y1))

        # First paint backgrounds and optional Excel gridlines. Painting text in a
        # second pass allows ordinary Excel text to flow into adjacent blank cells.
        for row, col, x0, y0, x1, y1 in cells:
            style = self.sheet.style(row, col)
            outline = GRID_COLOR if self.sheet.show_gridlines else ""
            self.body.create_rectangle(x0, y0, x1, y1, fill=style.fill, outline=outline, width=1)

        for row, col, x0, y0, x1, y1 in cells:
            style = self.sheet.style(row, col)
            text = self.sheet.text(row, col)
            if text:
                raw = self.sheet.value(row, col)
                horizontal = style.horizontal
                if horizontal == "general":
                    if isinstance(raw, (int, float, datetime, date, time)) and not isinstance(raw, bool):
                        horizontal = "right"
                    else:
                        horizontal = "left"

                padding = max(2, int(round(3 * self.zoom)))
                indent_px = max(0, int(round(style.indent * 9 * self.zoom)))
                if horizontal == "center":
                    tx, h_anchor = (x0 + x1) / 2, "center"
                elif horizontal == "right":
                    tx, h_anchor = x1 - padding - indent_px, "e"
                else:
                    tx, h_anchor = x0 + padding + indent_px, "w"

                if style.vertical == "top":
                    ty = y0 + padding
                    anchor = {"w": "nw", "center": "n", "e": "ne"}[h_anchor]
                elif style.vertical == "center":
                    ty = (y0 + y1) / 2
                    anchor = {"w": "w", "center": "center", "e": "e"}[h_anchor]
                else:
                    ty = y1 - padding
                    anchor = {"w": "sw", "center": "s", "e": "se"}[h_anchor]

                kwargs = {}
                if style.wrap:
                    kwargs["width"] = max(8, x1 - x0 - padding * 2)
                self.body.create_text(
                    tx,
                    ty,
                    text=text,
                    anchor=anchor,
                    fill=style.foreground,
                    font=self._font(style),
                    **kwargs,
                )

        # Explicit cell borders are drawn last so table outlines match Excel even
        # when worksheet gridlines are disabled.
        for row, col, x0, y0, x1, y1 in cells:
            style = self.sheet.style(row, col)
            self._draw_border_side(x0, y0, x1, y1, style.left, "left")
            self._draw_border_side(x0, y0, x1, y1, style.right, "right")
            self._draw_border_side(x0, y0, x1, y1, style.top, "top")
            self._draw_border_side(x0, y0, x1, y1, style.bottom, "bottom")

        self._draw_images(left, top, right, bottom)
        self._draw_selection()

    def _draw_images(self, left: float, top: float, right: float, bottom: float) -> None:
        for index, record in enumerate(self.sheet.images):
            x = self.columns.start(record.column) + int(round(record.x_offset * self.zoom))
            y = self.rows.start(record.row) + int(round(record.y_offset * self.zoom))
            if record.end_column is not None and record.end_row is not None:
                x2 = self.columns.start(record.end_column) + int(round(record.end_x_offset * self.zoom))
                y2 = self.rows.start(record.end_row) + int(round(record.end_y_offset * self.zoom))
                width = max(1, int(round(x2 - x)))
                height = max(1, int(round(y2 - y)))
            else:
                width = max(1, int(round(record.width * self.zoom)))
                height = max(1, int(round(record.height * self.zoom)))
            if x + width < left or x > right or y + height < top or y > bottom:
                continue
            try:
                image = Image.open(io.BytesIO(record.raw)).convert("RGBA")
                image = image.resize((width, height), Image.Resampling.LANCZOS)
                photo = ImageTk.PhotoImage(image)
                self._visible_images[index] = photo
                item = self.body.create_image(x, y, image=photo, anchor="nw", tags=(f"embedded_image_{index}", "embedded_image"))
                self.body.tag_raise(item)
                if self.selected_image == index:
                    self.body.create_rectangle(
                        x - 2,
                        y - 2,
                        x + width + 2,
                        y + height + 2,
                        outline=SELECTION_COLOR,
                        width=3,
                        tags=("image_selection",),
                    )
            except Exception:
                continue

    def _cell_from_event(self, event) -> tuple[int, int]:
        x = self.body.canvasx(event.x)
        y = self.body.canvasy(event.y)
        row, col = self.rows.index_at(y), self.columns.index_at(x)
        merged = self._cell_anchor(row, col)
        if merged is not None:
            min_row, min_col, _max_row, _max_col = merged
            return min_row, min_col
        return row, col

    def _image_from_event(self, event) -> Optional[int]:
        current = self.body.find_withtag("current")
        if not current:
            return None
        tags = self.body.gettags(current[0])
        for tag in tags:
            if tag.startswith("embedded_image_"):
                try:
                    return int(tag.rsplit("_", 1)[1])
                except ValueError:
                    return None
        return None

    def _mouse_down(self, event) -> str:
        self.body.focus_set()
        image_index = self._image_from_event(event)
        if image_index is not None:
            self.selected_image = image_index
            record = self.sheet.images[image_index]
            self.active_cell = (record.row, record.column)
            self.anchor_cell = self.active_cell
            self.status_callback(f"{self.sheet.title} | {record.label} at {_column_name(record.column)}{record.row} | Ctrl+C copies image")
            self.schedule_render()
            return "break"
        self.selected_image = None
        cell = self._cell_from_event(event)
        self.active_cell = cell
        self.anchor_cell = cell
        self._report_cell()
        self.schedule_render()
        return "break"

    def _mouse_drag(self, event) -> str:
        if self.selected_image is not None:
            return "break"
        self.active_cell = self._cell_from_event(event)
        self._report_cell()
        self.schedule_render()
        return "break"

    def _mouse_up(self, _event) -> str:
        return "break"

    def _right_click(self, event) -> str:
        image_index = self._image_from_event(event)
        if image_index is not None:
            self.selected_image = image_index
            self.schedule_render()
            menu = tk.Menu(self, tearoff=False)
            menu.add_command(label="Copy Image", command=self.copy_selected_image)
            menu.add_command(label="Save Image As...", command=self.save_selected_image)
            menu.tk_popup(event.x_root, event.y_root)
            return "break"
        cell = self._cell_from_event(event)
        self.active_cell = cell
        self.anchor_cell = cell
        self.selected_image = None
        self.schedule_render()
        menu = tk.Menu(self, tearoff=False)
        menu.add_command(label="Copy Cell / Range", command=self.copy_selection)
        menu.tk_popup(event.x_root, event.y_root)
        return "break"

    def _cell_in_selection(self, row: int, col: int) -> bool:
        r1, c1 = self.anchor_cell
        r2, c2 = self.active_cell
        return min(r1, r2) <= row <= max(r1, r2) and min(c1, c2) <= col <= max(c1, c2)

    def _draw_selection(self) -> None:
        if self.selected_image is not None:
            return
        r1, c1 = self.anchor_cell
        r2, c2 = self.active_cell
        r1, r2 = sorted((r1, r2))
        c1, c2 = sorted((c1, c2))
        x0 = self.columns.start(c1)
        x1 = self.columns.start(c2) + self.columns.size(c2)
        y0 = self.rows.start(r1)
        y1 = self.rows.start(r2) + self.rows.size(r2)
        if (r1, c1) == (r2, c2):
            merged = self._cell_anchor(r1, c1)
            if merged is not None:
                _min_row, _min_col, max_row, max_col = merged
                x1 = self.columns.start(max_col) + self.columns.size(max_col)
                y1 = self.rows.start(max_row) + self.rows.size(max_row)
        self.body.create_rectangle(x0, y0, x1, y1, outline=SELECTION_COLOR, width=3, tags=("selection",))
        self.body.tag_raise("selection")

    def _report_cell(self) -> None:
        row, col = self.active_cell
        text = self.sheet.text(row, col).replace("\n", " ")
        if len(text) > 160:
            text = text[:157] + "..."
        self.status_callback(f"{self.sheet.title} | {_column_name(col)}{row} = {text}")

    def _key_press(self, event) -> Optional[str]:
        key = event.keysym
        # Ctrl+Left / Ctrl+Right belong to the application's pane navigation.
        # Do not consume them as ordinary cell movement if they reach this layer.
        if key in {"Left", "Right"} and (int(getattr(event, "state", 0) or 0) & 0x0004):
            return None
        if key not in {"Left", "Right", "Up", "Down", "Home", "End", "Prior", "Next"}:
            return None
        row, col = self.active_cell
        if key == "Left":
            col = self.columns.next_visible(col, -1)
        elif key == "Right":
            col = self.columns.next_visible(col, 1)
        elif key == "Up":
            row = self.rows.next_visible(row, -1)
        elif key == "Down":
            row = self.rows.next_visible(row, 1)
        elif key == "Home":
            col = self.columns.first_visible()
        elif key == "End":
            col = self.columns.last_visible()
        elif key == "Prior":
            steps = max(1, int(self.body.winfo_height() / max(1, self.rows.default_size)))
            for _ in range(steps):
                new_row = self.rows.next_visible(row, -1)
                if new_row == row:
                    break
                row = new_row
        elif key == "Next":
            steps = max(1, int(self.body.winfo_height() / max(1, self.rows.default_size)))
            for _ in range(steps):
                new_row = self.rows.next_visible(row, 1)
                if new_row == row:
                    break
                row = new_row
        row = max(1, min(self.sheet.max_row, row))
        col = max(1, min(self.sheet.max_column, col))
        if not (event.state & 0x0001):
            self.anchor_cell = (row, col)
        self.active_cell = (row, col)
        self.selected_image = None
        self._ensure_visible(row, col)
        self._report_cell()
        self.schedule_render()
        return "break"

    def _ensure_visible(self, row: int, col: int) -> None:
        x0 = self.columns.start(col)
        x1 = x0 + self.columns.size(col)
        y0 = self.rows.start(row)
        y1 = y0 + self.rows.size(row)
        left = self.body.canvasx(0)
        top = self.body.canvasy(0)
        right = left + max(1, self.body.winfo_width())
        bottom = top + max(1, self.body.winfo_height())
        if x0 < left:
            self.body.xview_moveto(x0 / max(1, self.columns.total))
        elif x1 > right:
            self.body.xview_moveto(max(0.0, (x1 - self.body.winfo_width()) / max(1, self.columns.total)))
        if y0 < top:
            self.body.yview_moveto(y0 / max(1, self.rows.total))
        elif y1 > bottom:
            self.body.yview_moveto(max(0.0, (y1 - self.body.winfo_height()) / max(1, self.rows.total)))

    def copy_selection(self, _event=None) -> str:
        if self.selected_image is not None:
            self.copy_selected_image()
            return "break"
        r1, c1 = self.anchor_cell
        r2, c2 = self.active_cell
        r1, r2 = sorted((r1, r2))
        c1, c2 = sorted((c1, c2))
        rows = []
        for row in range(r1, r2 + 1):
            rows.append("\t".join(self.sheet.text(row, col) for col in range(c1, c2 + 1)))
        text = "\r\n".join(rows)
        self.clipboard_clear()
        self.clipboard_append(text)
        self.update_idletasks()
        self.status_callback(
            f"Copied cells {_column_name(c1)}{r1}:{_column_name(c2)}{r2} to clipboard"
        )
        return "break"

    def copy_selected_image(self) -> None:
        if self.selected_image is None:
            self.status_callback("Select an embedded image first.")
            return
        record = self.sheet.images[self.selected_image]
        try:
            image = Image.open(io.BytesIO(record.raw)).convert("RGBA")
            copy_pil_image_to_windows_clipboard(image)
            self.status_callback(f"Copied {record.label} to Windows clipboard")
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"Could not copy image:\n{exc}", parent=self.winfo_toplevel())

    def save_selected_image(self) -> None:
        if self.selected_image is None:
            return
        record = self.sheet.images[self.selected_image]
        extension = record.extension if record.extension.startswith(".") else ".png"
        target = filedialog.asksaveasfilename(
            parent=self.winfo_toplevel(),
            title="Save embedded image",
            defaultextension=extension,
            filetypes=[("Image", f"*{extension}"), ("PNG", "*.png"), ("All files", "*.*")],
        )
        if not target:
            return
        try:
            with Image.open(io.BytesIO(record.raw)) as image:
                image.save(target)
            self.status_callback(f"Saved image: {target}")
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"Could not save image:\n{exc}", parent=self.winfo_toplevel())

    def set_zoom(self, zoom: float) -> None:
        zoom = max(0.5, min(2.0, float(zoom)))
        if abs(zoom - self.zoom) < 0.001:
            return
        row, col = self.active_cell
        self.zoom = zoom
        self._font_cache.clear()
        self._build_metrics()
        self.body.configure(scrollregion=(0, 0, self.content_width, self.content_height))
        self.col_header.configure(scrollregion=(0, 0, self.content_width, 28))
        self.row_header.configure(scrollregion=(0, 0, 52, self.content_height))
        self._ensure_visible(row, col)
        self.render()


class ExcelViewerApp(tk.Tk):
    def __init__(self, initial_path: Optional[Path] = None) -> None:
        if os.name == "nt":
            try:
                import ctypes
                ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("jtl-sun.xViewer")
            except Exception:
                pass
        super().__init__()
        self.title(f"{APP_TITLE} {APP_VERSION}")
        try:
            icon_path = Path(__file__).resolve().parent / "assets" / "xviewer.ico"
            if icon_path.exists() and os.name == "nt":
                self.iconbitmap(default=str(icon_path))
        except Exception:
            pass
        self.geometry("1450x880")
        self.minsize(960, 620)
        self.model: Optional[WorkbookModel] = None
        self.path: Optional[Path] = None
        self.sheet_views: list[VirtualSheet] = []
        self.show_formulas = tk.BooleanVar(value=False)
        self.zoom_value = tk.StringVar(value="100%")
        self.status_text = tk.StringVar(value="Open an Excel workbook")
        self._loading = False
        self._build_ui()
        self.protocol("WM_DELETE_WINDOW", self._on_close)
        if initial_path is not None:
            self.after(50, lambda: self.open_path(initial_path))

    def _build_ui(self) -> None:
        toolbar = ttk.Frame(self, padding=(6, 5))
        toolbar.pack(fill="x")
        ttk.Button(toolbar, text="Open Excel", command=self.open_dialog).pack(side="left", padx=(0, 5))
        ttk.Button(toolbar, text="Copy", command=self.copy_current).pack(side="left", padx=3)
        ttk.Button(toolbar, text="Copy Image", command=self.copy_image_current).pack(side="left", padx=3)
        ttk.Button(toolbar, text="Open in Excel / Default App", command=self.open_default).pack(side="left", padx=(8, 3))
        ttk.Separator(toolbar, orient="vertical").pack(side="left", fill="y", padx=8)
        ttk.Label(toolbar, text="Zoom:").pack(side="left")
        zoom_box = ttk.Combobox(
            toolbar,
            textvariable=self.zoom_value,
            values=["50%", "75%", "90%", "100%", "110%", "125%", "150%", "175%", "200%"],
            width=7,
            state="readonly",
        )
        zoom_box.pack(side="left", padx=4)
        zoom_box.bind("<<ComboboxSelected>>", self._zoom_changed)
        ttk.Checkbutton(
            toolbar,
            text="Show formulas",
            variable=self.show_formulas,
            command=self.reload_current,
        ).pack(side="left", padx=(12, 3))

        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True)
        self.notebook.bind("<<NotebookTabChanged>>", lambda _event: self._sheet_changed())

        status = ttk.Label(self, textvariable=self.status_text, anchor="w", padding=(7, 4))
        status.pack(fill="x", side="bottom")

        self.bind("<Control-o>", lambda _event: self.open_dialog())
        self.bind("<Control-c>", lambda _event: self.copy_current())
        self.bind("<Control-C>", lambda _event: self.copy_current())
        self.bind("<Control-plus>", lambda _event: self.adjust_zoom(0.1))
        self.bind("<Control-equal>", lambda _event: self.adjust_zoom(0.1))
        self.bind("<Control-minus>", lambda _event: self.adjust_zoom(-0.1))
        self.bind("<Control-0>", lambda _event: self.set_zoom(1.0))

    def set_status(self, text: str) -> None:
        self.status_text.set(text)

    def open_dialog(self) -> None:
        if self._loading:
            return
        path = filedialog.askopenfilename(
            parent=self,
            title="Open Excel workbook",
            filetypes=[
                ("Excel workbooks", "*.xlsx *.xlsm *.xltx *.xltm *.xls"),
                ("Modern Excel", "*.xlsx *.xlsm *.xltx *.xltm"),
                ("Legacy Excel", "*.xls"),
                ("All files", "*.*"),
            ],
        )
        if path:
            self.open_path(Path(path))

    def open_path(self, path: Path) -> None:
        path = Path(path).expanduser()
        if path.suffix.lower() not in EXCEL_EXTENSIONS:
            messagebox.showwarning(APP_TITLE, f"Unsupported file type: {path.suffix}", parent=self)
            return
        if not path.exists():
            messagebox.showerror(APP_TITLE, f"File not found:\n{path}", parent=self)
            return
        if self._loading:
            return
        self._loading = True
        show_formulas = bool(self.show_formulas.get())
        self.set_status(f"Loading {path.name} ...")
        self.title(f"{APP_TITLE} {APP_VERSION} — Loading {path.name}")

        def worker() -> None:
            try:
                model = WorkbookModel(path, show_formulas=show_formulas)
            except Exception as exc:
                self.after(0, lambda: self._load_failed(path, exc))
                return
            self.after(0, lambda: self._load_complete(path, model))

        threading.Thread(target=worker, name="xExcel-Workbook-Loader", daemon=True).start()

    def _load_failed(self, path: Path, exc: Exception) -> None:
        self._loading = False
        self.set_status(f"Open failed: {path.name}")
        self.title(f"{APP_TITLE} {APP_VERSION}")
        messagebox.showerror(APP_TITLE, f"Could not open workbook:\n\n{path}\n\n{exc}", parent=self)

    def _load_complete(self, path: Path, model: WorkbookModel) -> None:
        old_model = self.model
        for tab in self.notebook.tabs():
            self.notebook.forget(tab)
        self.sheet_views.clear()
        self.model = model
        self.path = path
        for sheet in model.sheets:
            view = VirtualSheet(self.notebook, sheet, self.set_status)
            self.sheet_views.append(view)
            self.notebook.add(view, text=sheet.title)
        if old_model is not None:
            old_model.close()
        self._loading = False
        self.title(f"{APP_TITLE} {APP_VERSION} — {path.name}")
        image_count = sum(len(sheet.images) for sheet in model.sheets)
        self.set_status(
            f"{path} | {len(model.sheets)} sheet(s) | {image_count} embedded image(s) | "
            "Drag to select cells; Ctrl+C copies cells or the selected image"
        )
        if self.sheet_views:
            self.sheet_views[0].body.focus_set()

    def current_view(self) -> Optional[VirtualSheet]:
        if not self.sheet_views:
            return None
        selected = self.notebook.select()
        if not selected:
            return self.sheet_views[0]
        widget = self.nametowidget(selected)
        return widget if isinstance(widget, VirtualSheet) else None

    def copy_current(self) -> None:
        view = self.current_view()
        if view is not None:
            view.copy_selection()

    def copy_image_current(self) -> None:
        view = self.current_view()
        if view is not None:
            view.copy_selected_image()

    def open_default(self) -> None:
        if self.path is None:
            return
        try:
            os.startfile(str(self.path))  # type: ignore[attr-defined]
        except Exception as exc:
            messagebox.showerror(APP_TITLE, f"Could not open default application:\n{exc}", parent=self)

    def reload_current(self) -> None:
        if self.path is not None and not self._loading:
            self.open_path(self.path)

    def _zoom_changed(self, _event=None) -> None:
        text = self.zoom_value.get().strip().rstrip("%")
        try:
            self.set_zoom(float(text) / 100.0)
        except ValueError:
            self.zoom_value.set("100%")

    def set_zoom(self, zoom: float) -> None:
        zoom = max(0.5, min(2.0, zoom))
        self.zoom_value.set(f"{int(round(zoom * 100))}%")
        view = self.current_view()
        if view is not None:
            view.set_zoom(zoom)

    def adjust_zoom(self, delta: float) -> None:
        view = self.current_view()
        current = view.zoom if view is not None else 1.0
        self.set_zoom(current + delta)

    def _sheet_changed(self) -> None:
        view = self.current_view()
        if view is not None:
            view.body.focus_set()
            self.zoom_value.set(f"{int(round(view.zoom * 100))}%")
            view._report_cell()

    def _on_close(self) -> None:
        if self.model is not None:
            self.model.close()
        self.destroy()


def launch_excel_viewer(path: Optional[Path] = None) -> int:
    app = ExcelViewerApp(path)
    app.mainloop()
    return 0


def main(argv: Optional[Iterable[str]] = None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    initial = Path(args[0]) if args else None
    return launch_excel_viewer(initial)


if __name__ == "__main__":
    raise SystemExit(main())
